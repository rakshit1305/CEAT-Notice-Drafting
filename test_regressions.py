"""One test per defect found when the agent's four notices were reviewed.

Run:  python -m pytest -q tests/
No API key is needed — the model is simulated where the AI path is tested.
"""
import copy
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import pytest                                                    # noqa: E402

from core import analyse as A                                    # noqa: E402
from core import compose as C                                    # noqa: E402
from core import draft as D                                      # noqa: E402
from core import template as T                                   # noqa: E402
from core.validate import validate                               # noqa: E402
from fixtures import (BREACH, BREACH_FIXED, CONSUMER, CONSUMER_FIXED, RECOVERY,  # noqa: E402
                      RECOVERY_FIXED, S138, S138_FIXED)


def text(kind, case):
    return D.to_text(kind, copy.deepcopy(case))


def blockers(kind, case):
    return " | ".join(validate(kind, copy.deepcopy(case)).blockers)


def flags(kind, case):
    return " | ".join(f for _, f in validate(kind, copy.deepcopy(case)).flags)


# ============================================================ Section 138 ===
def test_s138_invoice_total_demand_is_blocked():
    b = blockers("s138", S138)
    assert "4,48,000" in b and "cheque" in b.lower()


def test_s138_no_part_payment_text_is_not_an_explanation():
    """'No further part-payment…' used to downgrade the mismatch to a warning."""
    assert not validate("s138", copy.deepcopy(S138)).ok


def test_s138_analysis_keeps_cheque_amount(monkeypatch):
    """The analysis step overwrote a correct ₹2,00,000 with the ₹4,48,000
    invoice total (_prefer_filtered_rows). Simulate the model and check."""
    case = {k: v for k, v in S138.items() if k not in ("amount", "part_payment")}
    model = A.Found(used_model=True, values={
        "amount": "200000",
        "part_payment": "No further part-payment has been made. Amount demanded is limited to INR 2,00,000/-.",
        "invoices": S138["invoices"], "cheques": S138["cheques"]})
    monkeypatch.setattr(A, "llm_ready", lambda: True)
    monkeypatch.setattr(A, "_call_text", lambda *a, **k: model)
    case["_raw_amt"] = "no part-payment; INR 2,00,000 demanded"
    from core.extract import Doc
    # the trigger: an attached invoice sheet whose total is not the cheque amount
    sheet = Doc(name="invoices.xlsx", kind="tables", digest="t1", tables=[{"sheet": "Sheet1", "rows": [
        ["Invoice No", "Invoice Date", "Customer", "Amount"],
        ["CEAT/CH/2026/0298", "28.05.2026", "Highway Auto Spares Private Limited", "448000"]]}])
    found = A.analyse("s138", case, [sheet])
    assert float(found.values["amount"]) == 200000
    assert "part_payment" not in found.values or not A._NO_PART.search(found.values["part_payment"])


def test_s138_no_drafting_note_paragraph():
    t = text("s138", S138_FIXED)
    assert "No further part-payment" not in t and "limited to" not in t


def test_s138_no_balance_when_nothing_paid():
    t = text("s138", S138_FIXED)
    assert "balance amount" not in t and "payment of the amount against the dishonoured cheque" in t


def test_s138_part_payment_is_stated_as_fact_with_balance():
    c = dict(S138_FIXED, part_paid="50000", part_paid_date="2026-06-25", amount="150000")
    t = text("s138", c)
    assert "paid a sum of INR 50,000/-" in t and "balance of INR 1,50,000/-" in t
    assert validate("s138", c).ok


def test_s138_raw_background_not_pasted():
    t = text("s138", S138)
    assert "3. Invoice No." not in t and "engaged in Dealer" not in t


def test_s138_director_address_not_doubled():
    t = text("s138", S138_FIXED)
    assert "Director, M/s Highway Auto Spares Private Limited, 14," not in t


def test_s138_corrected_passes_with_no_blanks():
    r = validate("s138", copy.deepcopy(S138_FIXED))
    assert r.ok and not r.blanks


# ============================================================== Recovery ===
def test_recovery_not_yet_due_invoices_blocked():
    b = blockers("recovery", RECOVERY)
    assert "not yet due" in b and "3,79,500" in b


def test_recovery_prior_question_not_assumed_no():
    assert "prior notices" in flags("recovery", RECOVERY)


def test_recovery_prior_notice_paragraph_and_overlap_flag():
    t = text("recovery", RECOVERY_FIXED)
    assert "separate notice dated 20.06.2026" in t
    assert "CEAT/CH/2026/0298 are in this statement of account" in flags("recovery", RECOVERY_FIXED)


def test_recovery_interest_start_and_jurisdiction():
    t = text("recovery", RECOVERY_FIXED)
    assert "8% p.a. from the due date of each invoice until realisation." in t
    assert "jurisdiction of the courts at Mumbai" in t
    assert "without a start date" in flags("recovery", RECOVERY)


def test_recovery_corrected_passes():
    r = validate("recovery", copy.deepcopy(RECOVERY_FIXED))
    assert r.ok and not r.blanks


# ============================================================== Consumer ===
def test_consumer_ceat_as_client_blocked():
    assert "CEAT's own" in blockers("consumer", CONSUMER)


def test_consumer_ceat_rejected_from_model_output():
    assert not A._plausible("client_name", "CEAT Limited")
    assert not A._plausible("client_address", "RPG House, 463 Dr. Annie Besant Road")


def test_consumer_subject_and_salutation_at_top():
    t = text("consumer", CONSUMER_FIXED)
    assert t.index("Sub: Reply to your Said Notice.") < t.index("We, CEAT Limited")
    assert t.index("Dear Sir/Madam,") < t.index("We, CEAT Limited")


def test_consumer_no_placeholder_for_missing_inspection():
    t = text("consumer", CONSUMER)
    assert "[● inspection finding]" not in t and "not eligible" not in t


def test_consumer_deny_then_admit_becomes_limited_admission():
    t = text("consumer", CONSUMER_FIXED)
    p2 = [l for l in t.split("\n") if l.startswith("With reference to para 2 ")][0]
    assert "admit only to the limited extent" in p2 and "deny the contents thereof" not in p2


def test_consumer_no_privity_block_not_used_when_dealer_authorised():
    t = text("consumer", CONSUMER)      # para 2 answer calls the dealer "authorised"
    assert "not purchased tyres from us or our authorized dealer" not in t


def test_consumer_contradiction_blocked_if_forced():
    c = dict(CONSUMER_FIXED, dealer_authorised="No")
    assert "Contradiction" in blockers("consumer", c)


def test_consumer_demands_not_denied_twice():
    t = text("consumer", CONSUMER_FIXED)
    assert "The demands raised in the Said Notice" not in t


def test_consumer_same_date_flagged_and_not_seeded():
    assert "same date" in flags("consumer", CONSUMER)


def test_consumer_uses_skill_wording_verbatim():
    t = text("consumer", CONSUMER_FIXED)
    assert "It also appears that your client has not properly apprised you of the facts." in t
    assert "on receipt we issue a claim receipt" in t
    assert "kindly note that we shall have no alternative" in t


def test_consumer_corrected_passes():
    r = validate("consumer", copy.deepcopy(CONSUMER_FIXED))
    assert r.ok and not r.blanks


# ================================================================ Breach ===
def test_breach_sentences_fitted_not_pasted():
    t = text("breach", BREACH)
    for bad in ("obliged to Timely", "in that Under", "constrained to If", "losses., at", ".."):
        assert bad not in t, bad
    assert "Clauses 9.1, 9.2 and 18.1" in t


def test_breach_amounts_indian_format():
    t = text("breach", BREACH)
    assert "INR 2,84,000/-" in t and "284000.00" not in t


def test_breach_proprietor_must_be_named():
    assert "proprietor" in blockers("breach", BREACH).lower()


def test_breach_unexplained_clause_flagged():
    assert "18.1" in flags("breach", BREACH)


def test_breach_corrected_passes():
    r = validate("breach", copy.deepcopy(BREACH_FIXED))
    assert r.ok and not r.blanks


# ======================================================= template source ===
def test_wording_comes_from_skill_file(monkeypatch):
    """Editing the skill must change the notice (the README promised this; the
    old code ignored the template text entirely)."""
    orig = T.paragraphs("s138")
    edited = tuple(p.replace("to our shock", "to our surprise") for p in orig)
    monkeypatch.setattr(T, "paragraphs", lambda kind: edited if kind == "s138" else orig)
    assert "to our surprise" in text("s138", S138_FIXED)


def test_missing_template_paragraph_reported_as_drift(monkeypatch):
    orig = T.paragraphs("s138")
    trimmed = tuple(p for p in orig if not p.startswith("That in the course of business"))
    monkeypatch.setattr(T, "paragraphs", lambda kind: trimmed if kind == "s138" else orig)
    assert "Template drift" in flags("s138", S138_FIXED)


# ============================================================ blanks gate ===
def test_blanks_reported():
    c = dict(S138_FIXED)
    c.pop("memo_date")
    c["cheques"] = [dict(S138_FIXED["cheques"][0], bank="")]
    r = validate("s138", c)
    assert any("drawn-on bank" in b for b in r.blanks)


def test_specimen_docx_is_marked():
    import io
    import docx
    raw = D.to_docx("s138", copy.deepcopy(S138_FIXED), specimen=True)
    first = docx.Document(io.BytesIO(raw)).paragraphs[0].text
    assert first.startswith("FILL-IN SPECIMEN")


# ======================================================== the AI paths ===
def test_ai_fit_accepted_when_faithful():
    case = copy.deepcopy(BREACH_FIXED)
    reply = json.dumps({"mode": "inline", "text": "make payment for all Products within 30 (thirty) days "
                                                  "of each invoice and pay interest at 8% per annum on "
                                                  "overdue amounts"})
    C.llm_fit_case("breach", case, chat=lambda m: reply)
    t = text("breach", case)
    assert "you were obliged to make payment for all Products within 30 (thirty) days" in t


def test_ai_fit_rejected_when_it_adds_a_fact():
    case = copy.deepcopy(BREACH_FIXED)
    reply = json.dumps({"mode": "inline", "text": "pay INR 9,99,999/- by 01.01.2027"})
    notes = C.llm_fit_case("breach", case, chat=lambda m: reply)
    assert any("rejected" in n for n in notes)
    assert "9,99,999" not in text("breach", case)


def test_ai_fit_goes_stale_when_answer_changes():
    case = copy.deepcopy(BREACH_FIXED)
    C.llm_fit_case("breach", case, chat=lambda m: json.dumps(
        {"mode": "inline", "text": "make payment within 30 (thirty) days"}))
    case["obligation"] = "maintain minimum stock of 400 tyres"
    assert "maintain minimum stock of 400 tyres" in text("breach", case)


def test_ai_review_flags_are_collected():
    reply = json.dumps({"issues": [{"severity": "high", "where": "para 10",
                                    "problem": "Demand differs from cheque amount"}]})
    out = C.ai_review("s138", "notice text", {"amount": "1"}, chat=lambda m: reply)
    assert out == [("crit", "AI read-through (para 10): Demand differs from cheque amount")]


def test_chat_retries_without_temperature(monkeypatch):
    """gpt-5.x rejects temperature=0; that used to kill every model call."""
    calls = []

    class Completions:
        def create(self, **kw):
            calls.append(kw)
            if "temperature" in kw:
                raise Exception("Unsupported value: 'temperature' does not support 0 with this model.")
            return "ok"

    class Client:
        chat = type("X", (), {"completions": Completions()})()

    monkeypatch.setattr("core.models.pick", lambda role: ("gpt-5.1", ""))
    assert A._chat(Client(), "text", temperature=0, messages=[]) == "ok"
    assert "temperature" not in calls[-1]


def test_openai_never_falls_back_to_arbitrary_model(monkeypatch):
    from core import models as M
    monkeypatch.setattr(M, "llm_settings", lambda: dict(key="k", base="", text="gpt-5.1", vision="gpt-5.1",
                                                        name="OpenAI-compatible"))
    monkeypatch.setattr(M, "available", lambda force=False: {"babbage-002", "gpt-4o"})
    assert M.pick("text")[0] == "gpt-4o"
    monkeypatch.setattr(M, "available", lambda force=False: {"babbage-002", "davinci-002"})
    assert M.pick("text")[0] == "gpt-5.1"


# ======================================================== freetext parse ===
@pytest.mark.parametrize("answer,paid,amount", [
    ("No further part-payment has been made. Amount demanded is limited to INR 2,00,000/-.", "0", "200000"),
    ("INR 50,000 received on 25.06.2026; balance INR 1,50,000 demanded.", "50000", "150000"),
])
def test_part_payment_answer_parsed(answer, paid, amount):
    from core.freetext import _parse
    v = _parse("s138", "amt", answer, {})
    assert v["part_paid"] == paid and float(v["amount"]) == float(amount)


def test_default_chat_path_through_fake_client(monkeypatch):
    """The real wiring: compose -> analyse._client/_chat -> client.chat.completions."""
    class Msg:
        content = json.dumps({"mode": "inline", "text": "terminate the Agreement for cause under Clause 19.1"})

    class Rsp:
        choices = [type("Ch", (), {"message": Msg()})()]

    class Completions:
        def create(self, **kw):
            assert kw["response_format"] == {"type": "json_object"}
            return Rsp()

    class Client:
        chat = type("X", (), {"completions": Completions()})()

    monkeypatch.setattr(A, "_client", lambda: (Client(), {}))
    monkeypatch.setattr("core.models.pick", lambda role: ("gpt-5.1", ""))
    case = copy.deepcopy(BREACH_FIXED)
    case["consequences"] = "terminate the Dealership Agreement for cause under Clause 19.1"
    C.llm_fit_case("breach", case)
    assert case["_fit"]["consequences"]["text"].startswith("terminate the Agreement")
    assert "constrained to terminate the Agreement for cause under Clause 19.1" in text("breach", case)
