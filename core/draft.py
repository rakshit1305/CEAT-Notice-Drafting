"""Assemble the notice from CEAT's approved wording, and write the .docx.

Fixed paragraphs are read from the skill's reference files at draft time (see
template.py); free-text answers are fitted into their sentences (compose.py)
rather than pasted. Anything unsupplied renders as [● label] in the preview —
and a notice with any [● …] left cannot be downloaded as a clean notice.
"""
from __future__ import annotations
import re
import io
from dataclasses import dataclass, field

from . import compose as C
from . import skill_loader as SK
from . import template as T
from .config import blank
from .schema import label, part_paid, rows
from .words import fmt_amount, fmt_date, house_words, inr, to_float, to_words

MODES_DEFAULT = "BY SPEED POST"


# --------------------------------------------------------------------------
# a tiny document model: paragraphs, numbered lists and tables
# --------------------------------------------------------------------------
@dataclass
class Block:
    kind: str                  # p | ol | list | table | stamp | addr | sig | subject | re
    text: str = ""
    items: list = field(default_factory=list)
    head: list = field(default_factory=list)
    rows: list = field(default_factory=list)
    bold: bool = False
    underline: bool = False


def V(case, key, lbl=None):
    v = case.get(key)
    if v is None or str(v).strip() == "":
        return blank(lbl or label(key).lower())
    return str(v).strip()


def D(case, key, lbl=None):
    v = case.get(key)
    s = fmt_date(v) if v else ""
    return s or blank(lbl or label(key).lower())


def M(case, key, lbl=None):
    v = to_float(case.get(key))
    return inr(v) if v is not None else blank(lbl or label(key).lower())


def W(case, key="amount", lbl="amount in words"):
    typed = str(case.get("amount_words") or "").strip()
    if typed:
        return house_words(typed)      # 'Rupees ... Only', however it was typed
    v = to_float(case.get(key))
    return to_words(v) if v is not None else blank(lbl)


def chq_facts(case, chq):
    """Per-cheque (no, returned-on, reason, memo date), each falling back to the
    single answer given for the whole notice. Uniform means one sentence will do."""
    rows = []
    for c in chq:
        rows.append((
            str(c.get("no", "")).strip(),
            str(c.get("dis", "")).strip() or str(case.get("dishonour_date", "")).strip(),
            str(c.get("reason", "")).strip() or str(case.get("dishonour_reason", "")).strip(),
            str(c.get("memo", "")).strip() or str(case.get("memo_date", "")).strip(),
        ))
    return rows, len({r[1:] for r in rows}) <= 1



def multi(case) -> bool:
    """Company+directors or Partnership+partners — both give numbered noticees."""
    t = str(case.get("noticee_type", "") or "")
    return t.startswith("Company") or t.startswith("Partnership")


def capacity(case) -> str:
    return "Partner" if str(case.get("noticee_type", "") or "").startswith("Partnership") else "Director"


# In these notices CEAT asserts nothing against anyone personally — it invokes a
# clause or gives information. A company's director is a contact, not a noticee;
# naming him as one implies he is personally answerable. A partnership is
# different: the partners are parties to the agreement and stay as noticees.
CONTACT_ONLY = ("fm", "price", "renewal")


def contact_line(case, kind: str = "") -> str:
    who = str(case.get("attn_name") or "").strip()
    desig = str(case.get("attn_desig") or "").strip()
    if not who:
        return ""
    return "Kind Attn.: " + who + (f", {desig}" if desig else "")


def noticee_block(case, kind: str = "") -> str:
    lines = []
    company_only = (kind in CONTACT_ONLY
                    and str(case.get("noticee_type", "")).startswith("Company"))
    if multi(case) and company_only:
        lines += [V(case, "noticee_name"), V(case, "noticee_address")]
        attn = contact_line(case, kind)
        named = [d.get("name") for d in rows(case, "directors") if str(d.get("name", "")).strip()]
        if not attn and named:
            d0 = rows(case, "directors")[0]
            cap0 = str(d0.get("capacity") or "").strip()
            attn = "Kind Attn.: " + str(d0.get("name")) + (f", {cap0}" if cap0 else "")
        if attn:
            lines += ["", attn]
    elif multi(case):
        lines += ["Noticee No. 1", V(case, "noticee_name"), V(case, "noticee_address")]
        for i, d in enumerate(rows(case, "directors"), start=2):
            if str(d.get("name", "")).strip():
                lines += ["", f"Noticee No. {i}", d["name"]]
                # The reference names the capacity and the company under each
                # director, so the block reads as a proper addressee.
                company = str(case.get("noticee_name") or "").strip()
                cap = str(d.get("capacity") or capacity(case)).strip()
                lines.append(f"{cap}, {company}" if company else cap)
                addr = str(d.get("address", "")).strip()
                if company:
                    # "Director, M/s X Pvt Ltd, 14 Industrial Area…" already has its
                    # capacity line above — printing it again doubled the line.
                    addr = re.sub(rf"^\s*(?:{re.escape(cap)}\s*,\s*)?{re.escape(company)}\s*,\s*", "",
                                  addr, flags=re.I)
                    addr = re.sub(rf"^\s*{re.escape(cap)}\s*,\s*", "", addr, flags=re.I)
                if addr:
                    lines.append(addr)
    else:
        lines.append(V(case, "noticee_name"))
        if str(case.get("firm_name", "")).strip():
            lines.append("Proprietor, " + case["firm_name"])
        lines.append(V(case, "noticee_address"))
        attn = contact_line(case, kind)
        if attn:
            lines += ["", attn]
    return "\n".join(lines)


def _join(xs: list[str]) -> str:
    xs = [x for x in xs if x]
    if len(xs) <= 1:
        return xs[0] if xs else ""
    return ", ".join(xs[:-1]) + " and " + xs[-1]


def _dir_phrase(dirs: list[str]) -> str:
    """Noticee numbering that follows the actual count of directors, so two
    directors do not silently render as the hard-coded 'Nos. 2 and 3'."""
    if not dirs:
        return blank("director names")
    nos = _join([str(i) for i in range(2, 2 + len(dirs))])
    word = "Noticee Nos." if len(dirs) > 1 else "Noticee No."
    return f"{word} {nos}, {_join(dirs)}"


def _discharge(inv, chq) -> str:
    """A cheque for less than the invoiced total is a PART discharge. Saying
    'in discharge' of a liability larger than the cheque misdescribes the debt."""
    it = sum(to_float(r.get("amt")) or 0 for r in (inv or []) if isinstance(r, dict))
    ct = sum(to_float(c.get("amt")) or 0 for c in (chq or []) if isinstance(c, dict))
    return "in part discharge" if (it and ct and ct < it - 0.5) else "in discharge"


def _presented(case) -> str:
    d = fmt_date(case.get("presented_date"))
    return f" on {d}" if d else " " + blank("date of presentation")


def _drawer_phrase(case) -> str:
    """Who signed the cheque — material where the drawer is a company."""
    if not multi(case):
        return ""
    return (", issued by Noticee No. 1 under the signature of its authorised "
            "signatory/director")


def _rate(case) -> str:
    """Bare rate. The answer box often already holds "8% p.a.", and the template
    appends "% p.a." itself — which printed "@ 8% p.a.% p.a."."""
    r = str(case.get("interest") or "8").strip()
    r = re.sub(r"\s*%?\s*(p\.?\s*a\.?|per\s+annum)\s*$", "", r, flags=re.I)
    return r.rstrip("%").strip() or "8"


def _authorised_dealer(case) -> bool:
    """Was the purchase from an authorised CEAT dealer?

    The no-privity defence asserts the client did NOT buy from us or our
    authorised dealer. Asserting that when the invoice says "Authorised Dealer –
    CEAT Limited" is a false statement in a signed reply, so the block switches.
    """
    blob = " ".join(str(case.get(k) or "") for k in
                    ("dealer", "client_relation", "product", "incoming", "extra_notes"))
    # the para-wise answers too: the old check missed "…Sterling Tyre World, an
    # authorised dealer…" sitting in the reply to para 2, and the reply then
    # asserted the opposite in block C
    blob += " " + " ".join(str(p.get("text") or "") for p in rows(case, "paras"))
    if re.search(r"\bOEM\b|original equipment", blob, re.I):
        return False
    hit = re.search(r"(?<!not an )(?<!not our )authori[sz]ed\s+(?:CEAT\s+)?dealer|CEAT\s+authori[sz]ed",
                    blob, re.I)
    return bool(hit)


def _as_sentence(rel: str, case) -> str:
    """The answer box holds a fragment — "appointed as an authorised dealer in
    2021..." — which cannot simply be appended. Give it a subject."""
    r = str(rel or "").strip().rstrip(".")
    if not r:
        return ""
    if re.match(r"^(appointed|engaged|supplied|introduced|nominated|granted)\b", r, re.I):
        r = "Noticee No. 1 was " + r[0].lower() + r[1:] if multi(case) else "You were " + r[0].lower() + r[1:]
    else:
        r = r[0].upper() + r[1:]
    return r + "."


def you(case) -> str:
    return ("you Noticees, jointly and severally,"
            if multi(case) else "you")


def sig_block(case) -> str:
    return ("Yours faithfully,\nFor CEAT Limited\n\n\n_________________\n"
            + V(case, "signatory_name") + "\n" + V(case, "signatory_desig"))


# --------------------------------------------------------------------------
# approved wording, looked up in the skill at draft time
# --------------------------------------------------------------------------
def TP(kind: str, anchor: str, fallback: str, values: dict | None = None, n: int | None = None) -> str:
    """The approved paragraph that opens with `anchor`, filled in.

    Taken from the skill's reference file so an edit there changes the notice.
    Only if the reference file no longer has it is the built-in copy used —
    and that is recorded, so the review notes can say the template drifted.
    """
    tpl = T.para(kind, anchor)
    if tpl is None:
        T.record(kind, f"paragraph beginning “{anchor}”")
        tpl = fallback
    try:
        return T.fill(tpl, values or {}, n)
    except T.MissingSlot as e:
        T.record(kind, f"paragraph beginning “{anchor}” now has a slot {{{{{e.args[0]}}}}} the app does not fill")
        return T.fill(fallback, values or {}, n)


def CL(heading: str, fallback: str, values: dict | None = None) -> str:
    """A block from clause-library.md, filled."""
    tpl = T.clause(heading)
    if tpl is None:
        T.record("clause-library", f"block “{heading}”")
        tpl = fallback
    try:
        return T.fill(tpl, values or {})
    except T.MissingSlot:
        return T.fill(fallback, values or {})


def BL(kind: str, label_: str, fallback: str, values: dict | None = None) -> str:
    """A labelled standard block from a reference file (consumer reply)."""
    tpl = T.block(kind, label_)
    if tpl is None:
        T.record(kind, f"block {label_}")
        tpl = fallback
    try:
        return T.fill(tpl, values or {})
    except T.MissingSlot:
        return T.fill(fallback, values or {})


def _clauses(s: str) -> str:
    """'Clause 9.1, 9.2 and 18.1' -> 'Clauses 9.1, 9.2 and 18.1'."""
    return re.sub(r"\bClause (\d[\d.()a-z]*(?:\s*(?:,|and|&)\s*(?:Clause\s+)?\d[\d.()a-z]*)+)",
                  r"Clauses \1", s)


def _final(s: str) -> str:
    return C.tidy(_clauses(s))


def _pct(row: dict) -> str:
    """The % change cell: as given, or worked out from the two prices."""
    p = str(row.get("pct") or "").strip()
    if p:
        return p if p.endswith("%") else p + "%"
    o, n = to_float(row.get("old")), to_float(row.get("nw"))
    if o and n:
        return f"{(n - o) / o * 100:+.1f}%"
    return "—"


def reg_city(ro: str) -> str:
    m = re.search(r"([A-Z][a-z]+)\s*[–-]\s*\d{6}", ro)
    return m.group(1) if m else "Mumbai"


def _prior_para(kind: str, case: dict, note) -> str:
    """The clause library's prior-notice block, when an earlier notice exists."""
    if not str(case.get("prior", "")).lower().startswith("y"):
        return ""
    pd_ = fmt_date(case.get("prior_date"))
    pref = C.strip_end(C.clean(case.get("prior_ref")))
    cur = C.strip_end(C.clean(case.get("current_instrument")))
    if not cur:
        cur = {"s138": "the dishonoured cheque(s) described above",
               "recovery": "the outstanding dues set out in the Statement of Account above",
               "breach": "the breach described above"}.get(kind, "the matter described above")
    if not pd_ or not pref:
        note("An earlier notice was reported, but its date or what it covered is missing, so the "
             "prior-notice paragraph could not be added. Give both to include it.")
        return ""
    return CL("Prior notice reference",
              "This notice pertains solely to {{CURRENT_INSTRUMENT_OR_BREACH}} and the amount/relief "
              "claimed herein does not include {{PRIOR_INSTRUMENT_OR_BREACH}}, which is the subject of "
              "our separate notice dated {{PRIOR_NOTICE_DATE}}. Nothing in this notice affects the "
              "Company's rights in respect of that separate matter, all of which are expressly reserved.",
              dict(CURRENT_INSTRUMENT_OR_BREACH=C.lower_first(cur),
                   PRIOR_INSTRUMENT_OR_BREACH=C.lower_first(pref), PRIOR_NOTICE_DATE=pd_))


def _jurisdiction(place: str) -> str:
    return CL("Jurisdiction",
              "Without prejudice to the Company's other rights, any legal proceedings arising out of or "
              "in connection with this notice shall be subject to the jurisdiction of the courts at "
              "{{JURISDICTION_PLACE}}.", dict(JURISDICTION_PLACE=place))


def dealer_status(case: dict) -> str:
    """'Yes' / 'No' / 'Not sure' — asked directly now; the old keyword guess is
    only a fallback, and it reads the para-wise answers too."""
    v = str(case.get("dealer_authorised") or "").strip()
    if v in ("Yes", "No", "Not sure"):
        return v
    return "Yes" if _authorised_dealer(case) else "Not sure"


ADMIT_PHRASE = "admit only to the limited extent stated below and deny the remainder"


def para_stance(p: dict) -> tuple[str, str, str]:
    """(stance phrase, cleaned body, note) for one para-wise answer.

    A body that admits part of the paragraph can no longer sit behind "we deny
    the contents thereof" — the stance becomes a limited admission and the
    lawyer is told. A body that merely restates the denial is trimmed.
    """
    n = str(p.get("n", "")).strip()
    stance = str(p.get("stance", "") or "").strip()
    # figures in a reply read as money too: "refund of 32000" -> "INR 32,000/-"
    body = C.money_full(C.clean(p.get("text", "")), words=False)
    body = re.sub(r"^with reference to (?:para|paragraph)\s*\d+[^,]*,\s*", "", body, flags=re.I)
    body = re.sub(r"^the contents of (?:para|paragraph)\s*\d+(?: of the (?:said |legal )?notice)?\s+"
                  r"(?:are|is)\s+(?:specifically |wholly |hereby |all )?denied\.?\s*", "", body, flags=re.I)
    body = C.cap_first(body) if body else ""
    note = ""
    if stance.startswith("Admit"):
        phrase = ADMIT_PHRASE
    elif stance.startswith("Not aware"):
        phrase = "are not aware of the averments and put your client to strict proof thereof"
    else:
        phrase = "deny the contents thereof"
        if body and C.ADMIT_RX.search(body):
            phrase = ADMIT_PHRASE
            note = (f"Para {n}: the answer admits part of the paragraph (“{C.ADMIT_RX.search(body).group(0)}”), "
                    "so “we deny the contents thereof” was changed to a limited admission. A limited "
                    "admission can be used in a later complaint — the lawyer must confirm it.")
    return phrase, body, note


# --------------------------------------------------------------------------
def build(kind: str, case: dict, notes: list | None = None) -> tuple[str, list[Block]]:
    ro = SK.reg_office()
    T.reset(kind)
    head: list[Block] = []
    body: list[Block] = []            # free paragraphs that come before the numbered list
    li: list = []
    tail: list[str] = []
    subject = ""
    opening = None
    collected: list[str] = [] if notes is None else notes

    def note(msg: str):
        if msg and msg not in collected:
            collected.append(msg)

    multi_ = multi(case)
    if (kind in CONTACT_ONLY and str(case.get("noticee_type", "")).startswith("Company")
            and [d for d in rows(case, "directors") if str(d.get("name", "")).strip()]):
        note("The named individual is shown as “Kind Attn.”, not as a noticee. This notice asserts "
             "nothing against anyone personally, and naming a director as a noticee implies he is "
             "personally answerable. Say so under ‘Anything else’ if you do intend to address him "
             "as a party.")
    if kind == "consumer":
        head.append(Block("addr", text="To,\n" + V(case, "advocate_name", "advocate") + ", Advocate\n"
                                       + V(case, "advocate_address", "advocate's address")))
        head.append(Block("p", text=D(case, "reply_date", "date of this reply")))
        head.append(Block("re", text="Re: Your Notice dated " + D(case, "notice_date", "date of the incoming notice")
                                     + " (“Said Notice”)"))
        subject = "Reply to your Said Notice."
    else:
        head.append(Block("p", text=D(case, "notice_date")))
        head.append(Block("stamp", text=case.get("mode") or MODES_DEFAULT, bold=True))
        head.append(Block("stamp", text="WITHOUT PREJUDICE", bold=True, underline=True))
        head.append(Block("addr", text="To,\n" + noticee_block(case, kind)))

    # ------------------------------------------------------------ s138 ---
    if kind == "s138":
        chq = rows(case, "cheques")
        inv = rows(case, "invoices")
        n_chq = max(len(chq), 1)
        many = len(chq) > 1
        paid = part_paid(case)
        subject = T.subject("s138") or "Legal Notice u/s 138 r/w Section 141 of the Negotiable Instruments Act, 1881."
        opening = TP("s138", "We, CEAT Limited, having our Registered Office",
                     f"We, CEAT Limited, having our Registered Office at {ro}, serve upon you the following notice:")
        li.append(TP("s138", "That we are a Public Limited Company",
                     "That we are a Public Limited Company duly registered under the provisions of the "
                     "Companies Act, 1956, having our registered office at the address mentioned above. We deal "
                     "in the business of manufacture and sale of Tyres, Tubes and Flaps (hereinafter referred "
                     "to as the \"Goods\")."))

        business, role, bnote = C.fit_business(case.get("business"), case)
        note(bnote)
        if not business:
            business = "the business of purchase and sale of the Goods"
            if multi_:
                note("No nature of business was given; the standard wording “the business of purchase "
                     "and sale of the Goods” was used. Correct it if the noticee does something else.")
        dirs = [d.get("name", "") for d in rows(case, "directors") if d.get("name")]
        if str(case.get("noticee_type", "")).startswith("Partnership"):
            li.append(f"That you Noticee No. 1, {V(case,'noticee_name')}, are a partnership firm "
                      f"registered under the Indian Partnership Act, 1932, engaged in {business}; and "
                      f"{_dir_phrase(dirs)}, are its partners responsible for its management and "
                      "day-to-day operations, and are jointly and severally liable for the obligations "
                      "of the said firm.")
        elif str(case.get("noticee_type", "")).startswith("Company"):
            li.append(f"That you Noticee No. 1, {V(case,'noticee_name')}, are a company registered in India "
                      f"engaged in {business}; and {_dir_phrase(dirs)}, are its directors responsible for "
                      "its management and day-to-day operations, and are liable under Section 141 of the "
                      "Negotiable Instruments Act, 1881.")
        else:
            li.append(CL("Noticee — individual",
                         "That you are the sole proprietor and the person in control and management of the "
                         "proprietorship concern {{FIRM_NAME}}, having your office and residential address "
                         "at {{ADDRESS}}, engaged in the business of purchase and sale of Goods.",
                         dict(FIRM_NAME=V(case, "firm_name", "firm name"),
                              ADDRESS=V(case, "noticee_address"))))

        bg, bgnote = C.fit_background(case.get("background"), case, multi_)
        note(bgnote)
        if not bg and role:
            bg, _ = C.fit_dealing(role, case, "Noticee No. 1" if multi_ else "you")
        if bg:
            li.append(bg)

        li.append(TP("s138", "That in the course of business",
                     "That in the course of business and against orders placed by you, the Company sold, "
                     "supplied and delivered the said Goods from time to time as per your requirements, which "
                     "were duly received, acknowledged and accepted by you without demur as to quality and/or "
                     "quantity."))
        if inv:
            lead = TP("s138", "That the Company supplied Goods vide",
                      "That the Company supplied Goods vide the following invoice(s), received by you "
                      "without complaint, against which you are liable to make payment:", n=len(inv))
            lead = lead.split(":")[0] + ":"
            li.append((lead, dict(head=["Invoice No.", "Invoice Date", "Amount (INR)"],
                                  rows=[[r.get("no") or blank("invoice no."),
                                         fmt_date(r.get("date")) or blank("invoice date"),
                                         fmt_amount(r.get("amt")) or blank("invoice amount")] for r in inv])))
        else:
            li.append("That the Company supplied Goods vide invoices received by you without complaint, "
                      "against which you are liable to make payment: " + blank("invoice particulars"))

        disc = _discharge(inv, chq)
        if many:
            li.append((f"That {disc} of your admitted liability you issued the following cheques:",
                       dict(head=["Cheque No.", "Cheque Date", "Amount (INR)", "Drawn on"],
                            rows=[[c.get("no", "—"), fmt_date(c.get("date")) or "—",
                                   fmt_amount(c.get("amt")) or "—", c.get("bank", "—")] for c in chq])))
        else:
            c = chq[0] if chq else {}
            amt_w = to_words(c.get("amt")) if to_float(c.get("amt")) else ""
            li.append(f"That {disc} of your admitted liability you issued cheque no. "
                      f"{c.get('no') or blank('cheque no.')} dated "
                      f"{fmt_date(c.get('date')) or blank('cheque date')} for INR "
                      f"{fmt_amount(c.get('amt')) or blank('cheque amount')}/-"
                      + (f" ({amt_w})" if amt_w else "")
                      + f" drawn on {c.get('bank') or blank('drawn-on bank')}{_drawer_phrase(case)}.")

        facts, uniform = chq_facts(case, chq)
        pres = fmt_date(case.get("presented_date"))
        if len(facts) > 1 and not uniform:
            li.append((f"That the aforesaid cheques were presented for clearance"
                       + (f" on {pres}" if pres else "")
                       + " and, to our shock, were returned unpaid by the drawee bank, as follows:",
                       dict(head=["Cheque No.", "Returned unpaid on", "Reason", "Bank memo dated"],
                            rows=[[n_ or "—", fmt_date(d) or blank("date of dishonour"),
                                   f"“{r}”" if r else blank("reason on memo"),
                                   fmt_date(m) or blank("bank memo date")] for n_, d, r, m in facts])))
        else:
            if facts:
                _, d0, r0, m0 = facts[0]
            else:
                d0 = str(case.get("dishonour_date", "")).strip()
                r0 = str(case.get("dishonour_reason", "")).strip()
                m0 = str(case.get("memo_date", "")).strip()
            t7 = TP("s138", "That the aforesaid cheque(s) was/were presented for clearance",
                    "That the aforesaid cheque(s) was/were presented for clearance and, to our shock, returned "
                    "unpaid on {{DISHONOUR_DATE}} for the reason \"{{DISHONOUR_REASON}}\" vide bank memo dated "
                    "{{DISHONOUR_MEMO_DATE}}.",
                    dict(DISHONOUR_DATE=fmt_date(d0) or blank("date of dishonour"),
                         DISHONOUR_REASON=r0 or blank("reason on memo"),
                         DISHONOUR_MEMO_DATE=fmt_date(m0) or blank("bank memo date")), n=n_chq)
            if pres:
                t7 = t7.replace("presented for clearance", f"presented for clearance on {pres}", 1)
            li.append(t7)

        # Part-payments: stated as facts with the balance, or not at all. The
        # answer box's own words are never printed — that is how "No further
        # part-payment has been made… the demand is limited to…" became para 8.
        if paid:
            tot = sum(to_float(c.get("amt")) or 0 for c in chq)
            pdate = fmt_date(case.get("part_paid_date"))
            bal = tot - paid
            li.append(f"That after the dishonour of the aforesaid {'cheques' if many else 'cheque'}, you "
                      f"have paid a sum of {inr(paid)} ({to_words(paid)})" + (f" on {pdate}" if pdate else "")
                      + f", and a balance of {inr(bal)} ({to_words(bal)}) remains payable against the "
                      f"dishonoured {'cheques' if many else 'cheque'}.")
        elif paid is None:
            note("A part-payment is mentioned but no amount is given, so no part-payment paragraph was "
                 "added. State the amount received and its date.")

        t9 = TP("s138", "That despite our repeated efforts",
                "That despite our repeated efforts to contact you for payment of the balance amount against "
                "the dishonoured cheque(s), you have failed and neglected to pay, demonstrating your disregard "
                "for the consequences of non-payment.", n=n_chq)
        if not paid:
            t9 = t9.replace("the balance amount", "the amount")   # nothing has been paid — no "balance"
        li.append(t9)
        li.append(TP("s138", "That your conduct reveals an intention tainted with fraud",
                     "That your conduct reveals an intention tainted with fraud. Prima facie you deliberately "
                     "issued the aforesaid cheque(s) with knowledge that it/they would be dishonoured, thereby "
                     "not only committing an offence under Section 138 of the Negotiable Instruments Act, 1881, "
                     "but also criminal breach of trust punishable under Section 316(2) of the Bharatiya Nyaya "
                     "Sanhita, 2023 (\"BNS\"), and cheating under Section 318(4) of the BNS.", n=n_chq))
        amt = to_float(case.get("amount"))
        t11 = TP("s138", "Hence, through this notice, you are called upon",
                 "Hence, through this notice, you are called upon to remit to the Company the entire legally "
                 "enforceable debt of INR {{AMOUNT_DEMANDED_FIGURES}} ({{AMOUNT_DEMANDED_WORDS}}) against the "
                 "dishonoured cheque(s), within a period of 15 (FIFTEEN) days from the date of receipt of this "
                 "notice, failing which the Company shall be constrained to initiate appropriate legal "
                 "proceedings against you, for which you shall be solely responsible for the costs and "
                 "consequences arising therefrom.",
                 dict(AMOUNT_DEMANDED_FIGURES=(fmt_amount(amt) + "/-") if amt else blank("amount demanded"),
                      AMOUNT_DEMANDED_WORDS=W(case)), n=n_chq)
        li.append(t11.replace("you are called upon", f"{you(case)} are called upon", 1))
        tail.append(CL("Retention note", "Please note that a copy of this notice has been retained by us "
                                         "for future reference."))

    # -------------------------------------------------------- recovery ---
    elif kind == "recovery":
        rate = _rate(case)
        amt = to_float(case.get("amount"))
        figs = (fmt_amount(amt) + "/-") if amt else blank("amount")
        words = W(case)
        subj = T.subject("recovery") or ("Demand notice for recovery of dues amounting to INR {{AMOUNT_FIGURES}} "
                                         "({{AMOUNT_WORDS}}) along with interest @ {{INTEREST_RATE}}% p.a.")
        subject = T.fill(subj, dict(AMOUNT_FIGURES=figs, AMOUNT_WORDS=words, INTEREST_RATE=rate))
        vals = dict(AMOUNT_FIGURES=figs, AMOUNT_WORDS=words, INTEREST_RATE=rate,
                    AS_ON_DATE=D(case, "as_on_date", "as-on date"))
        li.append(TP("recovery", "CEAT Limited (\"the Company\") is a Public Limited Company",
                     f"CEAT Limited (\"the Company\") is a Public Limited Company duly incorporated under the "
                     f"provisions of the Companies Act, 1956, having its Registered Office at {ro}."))
        li.append(TP("recovery", "The Company deals in the business",
                     "The Company deals in the business of manufacture and sale of Tyres, Tubes and Flaps "
                     "(hereinafter referred to as the \"Said Materials\")."))
        rel = case.get("relationship")
        fit = C.fitted(case, "relationship", rel)
        dirs = [d.get("name", "") for d in rows(case, "directors") if d.get("name")]
        if multi_ and dirs:
            kindword = "partners" if str(case.get("noticee_type", "")).startswith("Partnership") else "directors"
            firmword = ("a partnership firm registered under the Indian Partnership Act, 1932"
                        if kindword == "partners" else "a company registered in India")
            who = (f"That you Noticee No. 1, {V(case,'noticee_name')}, are {firmword} engaged in the "
                   f"business of purchase and sale of the Said Materials; and {_dir_phrase(dirs)}, "
                   f"are its {kindword} responsible for its management and day-to-day operations.")
            relsent = (C.end_stop(fit["text"]) if fit else C.fit_dealing(rel, case, "Noticee No. 1", prefix="")[0])
            li.append(who + (f" {relsent}" if relsent else ""))
        else:
            who = CL("Noticee — individual",
                     "That you are the sole proprietor and the person in control and management of the "
                     "proprietorship concern {{FIRM_NAME}}, having your office and residential address at "
                     "{{ADDRESS}}, engaged in the business of purchase and sale of Goods.",
                     dict(FIRM_NAME=V(case, "firm_name", "firm name"), ADDRESS=V(case, "noticee_address")))
            who = who.replace("purchase and sale of Goods", "purchase and sale of the Said Materials")
            relsent = (C.end_stop(fit["text"]) if fit else C.fit_dealing(rel, case, "you", prefix="")[0])
            li.append(who + (f" {relsent}" if relsent else ""))
        li.append(TP("recovery", "That in the course of business",
                     "That in the course of business and against orders placed by you, the Company sold, "
                     "supplied and delivered the Said Materials from time to time, duly received, acknowledged "
                     "and accepted by you without demur as to quality and/or quantity."))
        li.append(TP("recovery", "Towards the supply",
                     "Towards the supply, the Company raised various invoices and maintained a running and "
                     "continuous account of your transactions reflecting the amounts due and payable."))
        li.append(TP("recovery", "The Said Materials were received by you",
                     "The Said Materials were received by you in full satisfaction, and you were liable to pay "
                     "the invoice amounts, the Company having supplied on the assurance of payment within 30 "
                     "days from the date of invoice. However, you have failed to pay the outstanding invoices."))
        p7 = TP("recovery", "As per the Company's records",
                "As per the Company's records, as on {{AS_ON_DATE}} a total outstanding amount of INR "
                "{{AMOUNT_FIGURES}} ({{AMOUNT_WORDS}}) is pending from you towards the Said Materials supplied. "
                "The Statement of Account is as follows:", vals)
        soa = rows(case, "soa")
        if soa:
            total = sum(to_float(r.get("amt")) or 0 for r in soa)
            li.append((p7, dict(head=["SR No.", "Invoice / Reference", "Document Date", "Outstanding (INR)"],
                                rows=[[str(i), r.get("ref", "—"), fmt_date(r.get("date")) or "—",
                                       fmt_amount(r.get("amt")) or "—"] for i, r in enumerate(soa, 1)]
                                     + [["", "Total", "", fmt_amount(total)]])))
        else:
            li.append(re.sub(r"\s*The Statement of Account is as follows:\s*$", "", p7))
        li.append(TP("recovery", "Despite the Company's vigorous follow-ups",
                     "Despite the Company's vigorous follow-ups, you have failed to make payment as per the "
                     "agreed terms and conditions."))
        p9 = TP("recovery", "As per the agreed terms of the invoices",
                "As per the agreed terms of the invoices, you are liable and duty-bound to pay the outstanding "
                "amount of INR {{AMOUNT_FIGURES}} ({{AMOUNT_WORDS}}) as on {{AS_ON_DATE}}.", vals)
        li.append(p9.replace("you are liable", f"{you(case)} are liable", 1))
        if str(case.get("principal", "")).strip() or str(case.get("tax", "")).strip():
            li.append(CL("Amount breakup",
                         "The aforesaid outstanding amount of INR {{AMOUNT_FIGURES}} comprises a principal sum "
                         "of INR {{PRINCIPAL_AMOUNT}} and applicable GST/tax of INR {{TAX_AMOUNT}}, as reflected "
                         "in the invoices referred to above.",
                         dict(AMOUNT_FIGURES=figs,
                              PRINCIPAL_AMOUNT=(fmt_amount(case.get("principal")) + "/-")
                              if to_float(case.get("principal")) else blank("principal"),
                              TAX_AMOUNT=(fmt_amount(case.get("tax")) + "/-")
                              if to_float(case.get("tax")) else blank("tax component"))))
        li.append(TP("recovery", "You received the Said Materials",
                     "You received the Said Materials as per your requirement and yet failed to pay, "
                     "demonstrating a false and fraudulent intention from the beginning. Your acts and conduct "
                     "are patently illegal, untenable and contrary to settled principles of law."))
        p11 = TP("recovery", "In the aforesaid circumstances",
                 "In the aforesaid circumstances, the Company hereby calls upon you to pay the aforesaid sum of "
                 "INR {{AMOUNT_FIGURES}} ({{AMOUNT_WORDS}}) along with interest @ {{INTEREST_RATE}}% p.a.", vals)
        if multi_:
            p11 = p11.replace("calls upon you to pay", "calls upon you, jointly and severally, to pay", 1)
        frm = C.strip_end(C.clean(case.get("interest_from")))
        if frm:
            base11 = p11 if p11.endswith("p.a.") else p11.rstrip(" .")
            p11 = base11 + " " + (frm if frm.lower().startswith("from") else "from " + frm) \
                  + " until realisation."
        li.append(p11)
        li.append(TP("recovery", "Through your deliberate and wilful actions",
                     "Through your deliberate and wilful actions you have deceived and cheated the Company, "
                     "attracting the penal provisions of Sections 316(2) and 318(4) of the Bharatiya Nyaya "
                     "Sanhita, 2023, causing wrongful loss to the Company and wrongful gain to yourself."))
        li.append(TP("recovery", "In case you fail to pay the aforesaid amount",
                     "In case you fail to pay the aforesaid amount within 10 days from the date of receipt of "
                     "this notice, the Company will initiate appropriate legal proceedings, civil as well as "
                     "criminal, against you at your entire risk as to costs and consequences."))

    # -------------------------------------------------------- consumer ---
    elif kind == "consumer":
        opening_tpl = T.block("consumer", "**Opening / general denial:**") or (
            "We, CEAT Limited (\"Company\"), are in receipt of the Said Notice dated {{NOTICE_DATE}} issued on "
            "behalf of your client {{CLIENT_NAME}}, {{CLIENT_RELATION}}, residing at {{CLIENT_ADDRESS}}. At the "
            "outset, the Said Notice is factually baseless, patently incorrect, mala fide and unsustainable "
            "under law. Unless specifically admitted, nothing in the Said Notice shall be deemed admitted or "
            "accepted by us. It also appears that your client has not properly apprised you of the facts.")
        if T.block("consumer", "**Opening / general denial:**") is None:
            T.record("consumer", "block **Opening / general denial:**")
        # Optional details drop out cleanly instead of leaving a blank in the reply.
        if not str(case.get("client_relation") or "").strip():
            opening_tpl = opening_tpl.replace(", {{CLIENT_RELATION}}", "")
        if not str(case.get("client_address") or "").strip():
            opening_tpl = opening_tpl.replace(", residing at {{CLIENT_ADDRESS}}", "")
        body.append(Block("p", text=T.fill(opening_tpl, dict(
            NOTICE_DATE=D(case, "notice_date", "date of the incoming notice"),
            CLIENT_NAME=V(case, "client_name", "client's name"),
            CLIENT_RELATION=C.strip_end(C.clean(case.get("client_relation"))),
            CLIENT_ADDRESS=C.strip_end(C.clean(case.get("client_address")))))))
        if case.get("blk_a") is not False:
            body.append(Block("p", text="A) " + BL(
                "consumer", "**(A) Business / principal-to-principal:**",
                "We are engaged, inter alia, in manufacturing and marketing of tyres, tubes and flaps "
                "(\"Products\"). We sell our Products to our authorized dealers and to Original Equipment "
                "Manufacturers on a principal-to-principal basis. After such sale, we are not aware of onward "
                "sales by those dealers/OEMs to their customers/consumers.")))
        if case.get("blk_b") is not False:
            body.append(Block("p", text="B) " + BL(
                "consumer", "**(B) Warranty procedure:**",
                "The Products are subject to the warranty obligations detailed on our Company website, along "
                "with the procedure to claim warranty. Our procedure requires the tyre(s)/tube(s) under claim to "
                "be submitted to the dealer/OEM/us for inspection; on receipt we issue a claim receipt, our "
                "Technical Service Engineer examines the item, and the disposition is communicated to the "
                "consumer with a copy of the inspection report. For sizes that cannot be transported, our "
                "Engineer may inspect at the consumer's premises if desired.")))
        if case.get("blk_c") is not False:
            status = dealer_status(case)
            dealer = C.strip_end(C.clean(case.get("dealer"))) or blank("dealer / OEM")
            if status == "No":
                c_text = BL("consumer", "**(C) No privity of contract:**",
                            "It is an admitted fact that your client has not purchased tyres from us or our "
                            "authorized dealer. Your client purchased {{VEHICLE/PRODUCT}} from {{DEALER/OEM}}. We "
                            "have no contractual, commercial or legal relationship with your client, who is "
                            "admittedly not our direct customer. In the absence of privity of contract or other "
                            "legal nexus, we neither owe any obligation nor bear any responsibility towards your "
                            "client, and cannot be held liable for any representations or acts of {{DEALER/OEM}}. "
                            "Any such claims lie solely against the parties who dealt directly with your client.",
                            {"VEHICLE/PRODUCT": C.strip_end(C.clean(case.get("product"))) or blank("vehicle / product"),
                             "DEALER/OEM": dealer})
            elif status == "Yes":
                # The purchase WAS from an authorised dealer. The approved block
                # would assert the opposite — a false statement in a signed reply.
                c_text = ("We note that the purchase relied upon was made from " + dealer + ", which describes "
                          "itself as an authorised dealer of the Company. That invoice is the dealer's own "
                          "document, and the sale recorded therein is a sale by the dealer to your client on a "
                          "principal-to-principal basis. Our obligations, if any, are confined to the limited "
                          "warranty referred to above and to the procedure prescribed thereunder.")
                note("The purchase was from an authorised CEAT dealer, so the approved no-privity block "
                     "(which says the opposite) was replaced with a limited-warranty paragraph. That "
                     "paragraph is not CEAT-approved wording — legal must confirm it.")
            else:
                c_text = ("Your client has not purchased any Product directly from us. The purchase relied upon "
                          "was made from " + dealer + ", and the sale recorded in that dealer's invoice is a sale "
                          "by the dealer to your client on the dealer's own account. Our obligations, if any, are "
                          "confined to the limited warranty referred to above and the procedure prescribed "
                          "thereunder. Our right to verify the said invoice and the status of the said dealer "
                          "from our records is expressly reserved.")
                note("It is not confirmed whether the dealer is an authorised CEAT dealer, so the approved "
                     "no-privity block was NOT used (it would assert they are not). A cautious paragraph "
                     "that reserves the point was used instead — answer the dealer question to use the "
                     "approved block, and have legal confirm the wording.")
            body.append(Block("p", text="C) " + c_text))

        claim = fmt_date(case.get("claim_date"))
        insp = C.strip_end(C.clean(case.get("inspection")))
        if insp:
            t = ("For the record: " + (f"we received the warranty claim on {claim}; " if claim else "")
                 + f"on inspection the item was found {C.lower_first(insp)}, and it is accordingly not "
                   "eligible under our warranty policy")
            if str(case.get("rejection", "")).strip():
                t += f", which disposition was communicated to your client on {fmt_date(case['rejection'])}"
            body.append(Block("p", text=t + "."))
        elif claim:
            body.append(Block("p", text=f"For the record: we received the warranty claim on {claim}."))
            note("A warranty-claim date is recorded but no inspection finding, so the reply states only that "
                 "the claim was received — it asserts no outcome. Add the finding if there was one.")
        body.append(Block("p", text="Now, we respond para-wise to your Said Notice as follows:"))

        _paras = rows(case, "paras")
        answered_demands = False
        i = 0
        while i < len(_paras):
            p = _paras[i]
            stance = str(p.get("stance", "") or "").strip()
            phrase, ptext, pnote = para_stance(p)
            note(pnote)
            if re.search(r"\bdemand", ptext, re.I):
                answered_demands = True
            if not stance and not ptext:
                body.append(Block("p", text=(f"With reference to para {p.get('n')} of the Said Notice, "
                                             + blank("stance — admit / deny / not aware, and CEAT's answer"))))
                i += 1
                continue
            if not stance and ptext:
                body.append(Block("p", text=f"With reference to para {p.get('n')} of the Said Notice, "
                                            + blank("stance — admit / deny / not aware") + f" {ptext}"))
                i += 1
                continue
            if ptext:
                body.append(Block("p", text=f"With reference to para {p.get('n')} of the Said Notice, "
                                            f"we {phrase}. {ptext}"))
                i += 1
                continue
            j = i
            while (j + 1 < len(_paras)
                   and str(_paras[j + 1].get("stance", "") or "").strip() == stance
                   and not str(_paras[j + 1].get("text", "") or "").strip()):
                j += 1
            ref = (f"paras {_paras[i].get('n')} to {_paras[j].get('n')}" if j > i else f"para {p.get('n')}")
            body.append(Block("p", text=f"With reference to {ref} of the Said Notice, we {phrase}."))
            i = j + 1
        if not _paras:
            body.append(Block("p", text=blank("para-wise responses — one per numbered paragraph")))

        dem = [x for x in rows(case, "demands") if str(x.get("d", "")).strip()]
        denied = [C.strip_end(x["d"]) for x in dem if not str(x.get("resp", "")).startswith("Conceded")]
        if denied and not answered_demands:
            if len(denied) == 1:
                body.append(Block("p", text="The demand raised in the Said Notice, namely "
                                            + C.lower_first(C.money_full(denied[0], words=False))
                                            + ", is wholly untenable and is "
                                            "hereby expressly denied and rejected."))
            else:
                # One long sentence listing four demands is hard to read and
                # harder to answer paragraph by paragraph later.
                body.append(Block("p", text="The demands raised in the Said Notice, namely:"))
                body.append(Block("list", items=[_final(x) for x in C.punctuate(
                    [C.cap_first(C.money_full(d, words=False)) for d in denied], last=",")]))
                body.append(Block("p", text="are each wholly untenable and are hereby expressly denied "
                                            "and rejected."))
        elif denied:
            note("The demands are already answered in the para-wise reply, so the separate “demands are "
                 "denied” paragraph was not repeated.")
        for x in dem:
            if str(x.get("resp", "")).startswith("Conceded"):
                body.append(Block("p", text=f"In respect of your client's demand for {C.strip_end(x['d'])}, "
                                            f"{str(x.get('note','')).strip() or 'we shall act as stated separately'}."))
        body.append(Block("p", text=BL(
            "consumer", "**Closing:**",
            "Under such circumstances, notwithstanding the above, if your client initiates proceedings against "
            "us, kindly note that we shall have no alternative but to defend the same at your client's entire "
            "risk as to the costs and consequences thereof.")))

    # -------------------------------------------- contract notice types ---
    # These four (termination, renewal, force majeure, price) are NON-STANDARD in
    # the skill — no CEAT sample exists yet — so their wording comes from the
    # reference file's own template and every answer is fitted into it rather
    # than pasted mid-sentence.
    else:
        agr = V(case, "agreement_name", "agreement title")
        dt = D(case, "agreement_date")
        cl = V(case, "clause_no", "clause")
        if kind == "breach":
            av = dict(AGREEMENT_NAME=agr, AGREEMENT_DATE=dt, CLAUSE_NO=V(case, "clause_no", "clause"),
                      CURE_PERIOD=V(case, "cure_period", "cure period"))
            subject = T.fill(T.subject("breach") or "Notice of Breach of {{AGREEMENT_NAME}} dated "
                             "{{AGREEMENT_DATE}} — call to cure.", av)
            li.append(TP("breach", "CEAT Limited (\"the Company\") is a Public Limited Company",
                         f"CEAT Limited (\"the Company\") is a Public Limited Company incorporated under the "
                         f"Companies Act, 1956, having its Registered Office at {ro}, engaged in the manufacture "
                         f"and sale of Tyres, Tubes and Flaps."))
            li.append(TP("breach", "The Company and you entered into",
                         "The Company and you entered into {{AGREEMENT_NAME}} dated {{AGREEMENT_DATE}} (\"the "
                         "Agreement\").", av))
            # The three answers below used to be pasted into the middle of the
            # approved sentences verbatim ("obliged to Timely payment…",
            # "constrained to If you fail…"). They are now fitted to them.
            p3 = T.para("breach", "In terms of Clause")
            if p3 is None:
                T.record("breach", "paragraph beginning “In terms of Clause”")
            p3 = (p3 or "In terms of Clause {{CLAUSE_NO}} of the Agreement, you were obliged to {{OBLIGATION}}.")
            p3 = p3.replace("{{CLAUSE_NO}}", av["CLAUSE_NO"])
            ob, onote = C.fit_obligation(case.get("obligation"), case, p3)
            note(onote)
            li.append(ob or p3.replace("{{OBLIGATION}}", blank("the obligation")))
            p4 = T.para("breach", "You are in breach of the said obligation")
            if p4 is None:
                T.record("breach", "paragraph beginning “You are in breach of the said obligation”")
            fx, fnote = C.fit_breach_facts(case.get("breach_facts"), case, p4)
            note(fnote)
            li.append(fx or (p4 or "You are in breach of the said obligation, in that {{BREACH_FACTS}}. "
                             "Despite the Company's follow-ups, the breach subsists.").replace(
                "{{BREACH_FACTS}}", blank("particulars of the breach")))
            li.append(TP("breach", "You are hereby called upon to remedy",
                         "You are hereby called upon to remedy/cure the aforesaid breach within {{CURE_PERIOD}} "
                         "from the date of receipt of this notice.", av))
            p6 = T.para("breach", "Should you fail to cure the breach")
            if p6 is None:
                T.record("breach", "paragraph beginning “Should you fail to cure the breach”")
            cq, cnote = C.fit_consequences(case.get("consequences"), case, p6)
            note(cnote)
            li.append(cq or (p6 or "Should you fail to cure the breach within the said period, the Company "
                             "shall be constrained to {{CONSEQUENCES}}, at your entire risk as to costs and "
                             "consequences.").replace("{{CONSEQUENCES}}", blank("consequences")))
            li.append(TP("breach", "This notice is issued without prejudice",
                         "This notice is issued without prejudice to the Company's rights and remedies, all of "
                         "which are expressly reserved."))
        elif kind == "termination":
            intro = TP("termination", "CEAT Limited (\"the Company\") is a Public Limited Company",
                       f"CEAT Limited (“the Company”) is a Public Limited Company incorporated under the "
                       f"Companies Act, 1956, having its Registered Office at {ro}, engaged in the "
                       f"manufacture and sale of Tyres, Tubes and Flaps.")
            li.append(intro)
            subject = f"Notice of Termination of {agr} dated {dt}."
            li.append(TP("termination", "The Company and you entered into",
                         "The Company and you entered into {{AGREEMENT_NAME}} dated {{AGREEMENT_DATE}} "
                         "(\"the Agreement\"), governing {{SUBJECT_OF_AGREEMENT}}.",
                         dict(AGREEMENT_NAME=agr, AGREEMENT_DATE=dt,
                              SUBJECT_OF_AGREEMENT=C.fit_inline(case.get("subject_of"))
                              or blank("what the agreement governs"))))
            g = str(case.get("ground", ""))
            if g == "Convenience":
                li.append(f"In terms of Clause {cl}, the Company is entitled to terminate the Agreement "
                          f"by giving {C.fit_inline(case.get('notice_period')) or blank('notice period')} "
                          "notice, and has elected to do so.")
            elif g == "Expiry":
                li.append(f"In terms of Clause {cl}, the Agreement stands terminated upon expiry of its term.")
            else:
                # The obligation clause is its own field: citing the termination
                # clause as the source of the obligation (12.1(a) instead of 7.2)
                # misstates the agreement.
                ocl = str(case.get("obligation_clause") or "").strip()
                if not ocl:
                    ocl = str(case.get("clause_no") or "").strip()
                    if ocl:
                        note(f"No separate clause was given for the obligation, so the termination clause "
                             f"({ocl}) is cited for both. If the obligation sits in a different clause, "
                             "say which — citing the termination clause as the source of the obligation "
                             "misstates the agreement.")
                lead, subs = C.fit_required_to(case.get("obligation"), case, ocl or blank("clause"))
                li.append((lead, dict(sub=subs)) if subs else
                          (lead or f"In terms of Clause {ocl or blank('clause')}, you were required to "
                           + blank("the obligation")))
                facts = C.fit_failed_as(case.get("facts"), case)
                cure = ""
                if str(case.get("cure_given_date", "")).strip():
                    cure = (f" This is despite the cure period given on "
                            f"{fmt_date(case['cure_given_date'])}")
                    if str(case.get("cure_lapsed_date", "")).strip():
                        cure += f", which lapsed on {fmt_date(case['cure_lapsed_date'])}"
                    cure += "."
                li.append((facts or "You have failed to do so, as " + blank("what happened")) + cure)
            li.append(TP("termination", "Accordingly, the Company hereby gives notice of termination",
                         "Accordingly, the Company hereby gives notice of termination of the Agreement "
                         "under Clause {{CLAUSE_NO}}, with effect from {{EFFECTIVE_DATE}}.",
                         dict(CLAUSE_NO=cl, EFFECTIVE_DATE=D(case, "effective_date"))))
            lead = "Upon termination, you are called upon"
            wd_raw = str(case.get("wind_down") or "")
            dues_txt = str(case.get("dues") or "").strip()
            # Don't lead with the dues when the wind-down list already says to pay
            # them — the old version demanded the same money twice in one sentence.
            dues_shown = bool(dues_txt) and not re.search(
                re.escape(re.sub(r"[^\d]", "", dues_txt)[:6]), re.sub(r"[^\d]", "", wd_raw))
            if dues_txt and dues_shown:
                lead += f" to clear all outstanding dues of {M(case,'dues')} and"
            wlead, wsubs = C.fit_called_upon(case.get("wind_down"), case, "wind_down", lead + " to")
            if wsubs:
                li.append((wlead, dict(sub=wsubs)))
            else:
                li.append(wlead or lead + " to " + blank("wind-down obligations") + ".")
            li.append(TP("termination", "This termination is without prejudice",
                         "This termination is without prejudice to the Company's rights and remedies, all "
                         "of which are expressly reserved, including for recovery of dues and losses."))

        elif kind == "renewal":
            li.append(f"CEAT Limited (“the Company”) is a Public Limited Company incorporated under the "
                      f"Companies Act, 1956, having its Registered Office at {ro}, engaged in the "
                      f"manufacture and sale of Tyres, Tubes and Flaps.")
            renew = case.get("intent") != "Do not renew"
            subject = f"Notice of {'Renewal' if renew else 'Non-Renewal'} of {agr} dated {dt}."
            li.append(TP("renewal", "CEAT Limited (\"the Company\") and you are parties to",
                         "CEAT Limited (\"the Company\") and you are parties to {{AGREEMENT_NAME}} dated "
                         "{{AGREEMENT_DATE}} (\"the Agreement\"), which is due to expire on {{EXPIRY_DATE}}.",
                         dict(AGREEMENT_NAME=agr, AGREEMENT_DATE=dt,
                              EXPIRY_DATE=D(case, "expiry_date"))).replace(
                "CEAT Limited (“the Company”) and you", "The Company and you", 1))
            if renew:
                term = C.fit_inline(case.get("renewal_term")) or blank("renewal term")
                has_terms = bool(str(case.get("revised_terms") or "").strip())
                lead = (f"In terms of Clause {cl}, the Company hereby conveys its intention to renew the "
                        f"Agreement for a further term of {term}"
                        + (" on the following revised terms" if has_terms else " on the existing terms"))
                tlead, tsubs = C.fit_terms(case.get("revised_terms"), case, lead)
                li.append((tlead, dict(sub=tsubs)) if tsubs else tlead)
                by = D(case, "confirm_by", "confirm-by date")
                li.append(f"Kindly confirm your acceptance of the aforesaid terms in writing on or before "
                          f"{by}. The renewal is offered on the terms stated above and is open for "
                          f"acceptance until that date; if your written acceptance is not received by then, "
                          f"the Agreement shall expire on {D(case,'expiry_date')} in accordance with its "
                          "own terms, without further notice.")
            else:
                li.append(TP("renewal", "In terms of Clause {{CLAUSE_NO}}, the Company hereby gives notice "
                             "that it does NOT",
                             "In terms of Clause {{CLAUSE_NO}}, the Company hereby gives notice that it does "
                             "NOT intend to renew the Agreement, which shall accordingly stand expired on "
                             "{{EXPIRY_DATE}} without further notice.",
                             dict(CLAUSE_NO=cl, EXPIRY_DATE=D(case, "expiry_date"))))
                lead = "Upon expiry, you are called upon"
                wd_raw = str(case.get("wind_down") or "")
                dues_txt = str(case.get("dues") or "").strip()
                if dues_txt and not re.search(re.escape(re.sub(r"[^\d]", "", dues_txt)[:6]),
                                              re.sub(r"[^\d]", "", wd_raw)):
                    lead += f" to clear outstanding dues of {M(case,'dues')} and"
                wlead, wsubs = C.fit_called_upon(case.get("wind_down"), case, "wind_down", lead + " to")
                if wsubs:
                    li.append((wlead, dict(sub=wsubs)))
                elif wlead:
                    li.append(wlead)
            li.append("This notice is without prejudice to the Company’s rights and remedies, all of "
                      "which are expressly reserved.")

        elif kind == "fm":
            li.append(f"CEAT Limited (“the Company”) is a Public Limited Company incorporated under the "
                      f"Companies Act, 1956, having its Registered Office at {ro}, engaged in the "
                      f"manufacture and sale of Tyres, Tubes and Flaps.")
            subject = f"Notice invoking Force Majeure under {agr} dated {dt}."
            li.append(TP("fm", "The Company and you are parties to",
                         "The Company and you are parties to {{AGREEMENT_NAME}} dated {{AGREEMENT_DATE}} "
                         "(\"the Agreement\"), which contains a force majeure provision at Clause "
                         "{{CLAUSE_NO}}.", dict(AGREEMENT_NAME=agr, AGREEMENT_DATE=dt, CLAUSE_NO=cl)))
            li.append(C.fit_event(case.get("fm_event"), case, D(case, "fm_date", "event date"), cl)
                      or (f"On {D(case,'fm_date','event date')}, " + blank("the force majeure event")
                          + f" has occurred, being an event beyond the Company’s reasonable control "
                            f"within the meaning of Clause {cl}."))
            aff = C.fit_affected(case.get("affected"), case)
            imp = C.fit_impact(case.get("impact"), case)
            qty = [r for r in rows(case, "quantities") if str(r.get("period", "")).strip()]
            aff_txt = ((aff or "As a direct consequence, the Company is prevented or delayed from "
                        "performing " + blank("affected obligations") + ".")
                       + (" " + imp if imp else " The expected impact and duration is "
                          + blank("impact / duration") + "."))
            if qty:
                # The affected deliveries as a table: the same figures in prose are
                # far easier for the other side to dispute later.
                tot = {k: sum(to_float(r.get(k)) or 0 for r in qty) for k in ("sched", "done", "pend")}
                li.append((C.strip_end(aff_txt.split(". ")[0]) + ", in respect of the following "
                           "quantities:",
                           dict(head=["Month / period", "Scheduled", "Delivered", "Pending", "Due date"],
                                rows=[[r.get("period", "—"), r.get("sched") or "—", r.get("done") or "—",
                                       r.get("pend") or "—", fmt_date(r.get("due")) or r.get("due") or "—"]
                                      for r in qty]
                                     + [["Total", f"{tot['sched']:g}", f"{tot['done']:g}",
                                         f"{tot['pend']:g}", ""]])))
                rest = ". ".join(aff_txt.split(". ")[1:]).strip()
                if rest:
                    li.append(rest)
            else:
                li.append(aff_txt)
            mlead, msubs = C.fit_called_upon(
                case.get("mitigation"), case, "mitigation",
                "The Company is taking the following steps to mitigate the effect of the said event")
            if msubs:
                li.append((mlead, dict(sub=msubs)))
            else:
                li.append(mlead or "The Company is taking the following steps to mitigate the effect of "
                                   "the said event: " + blank("mitigation steps") + ".")
            rlead, rsubs = C.fit_relief(case.get("relief"), case, cl)
            keep = " The Company will keep you informed and resume performance as soon as reasonably " \
                   "practicable."
            if rsubs:
                li.append((rlead, dict(sub=rsubs)))
                li.append(keep.strip())
            else:
                li.append(rlead + keep)
            li.append("This notice is without prejudice to the Company’s rights and remedies under the "
                      "Agreement and in law, all of which are expressly reserved.")

        elif kind == "price":
            subject = f"Notice of Price Revision effective {D(case,'effective_date')}."
            basis = (f"Clause {case['clause_no']}" if str(case.get("clause_no", "")).strip()
                     else "the Company’s right to revise prices")
            li.append(f"CEAT Limited (“the Company”) supplies its Tyres, Tubes and Flaps (“Goods”) to you "
                      f"under {agr}" + (f" dated {fmt_date(case['agreement_date'])}"
                                        if str(case.get("agreement_date", "")).strip() else "") + ".")
            reason = C.fit_inline(C.house_self(str(case.get("reason") or ""))) or blank("reason for the revision")
            pct = str(case.get("price_change_pct") or "").strip()
            li.append(f"On account of {reason}, and in terms of {basis}, the Company hereby gives notice of "
                      f"a revision in the prices of the Goods"
                      + (f", by {pct}%," if pct and not rows(case, "prices") else "")
                      + f" with effect from {D(case,'effective_date')}.")
            prows = [r for r in rows(case, "prices") if str(r.get("sku", "")).strip()]
            if prows:
                li.append(("The revised prices are as follows:",
                           dict(head=["Product / SKU", "Existing price (INR)", "Revised price (INR)",
                                      "% change"],
                                rows=[[r.get("sku"), fmt_amount(r.get("old")) or "—",
                                       fmt_amount(r.get("nw")) or "—",
                                       _pct(r)] for r in prows])))
            elif pct:
                li.append(f"The revision is a uniform {pct}% across the products supplied under the said "
                          "arrangement; the revised price of each product is its existing price so revised.")
                note(f"Only an overall percentage ({pct}%) was given, so the notice states the percentage "
                     "and no size-wise table. A product-wise table is clearer and harder to dispute — add "
                     "one if you have it.")
            else:
                li.append("The revised prices are as follows: " + blank("price table, or the overall % change"))
            pre = C.fit_inline(C.house_self(str(case.get("pre_orders") or "")))
            li.append(f"Orders placed and accepted before {D(case,'effective_date')} shall be "
                      + (pre or blank("treatment of pre-existing orders"))
                      + f". All orders accepted on or after {D(case,'effective_date')} shall be at the "
                        "revised prices.")
            li.append("All other terms and conditions of supply remain unchanged. This notice is without "
                      "prejudice to the Company’s rights, all of which are expressly reserved.")

    # ---------------------------------------------------------- assemble ---
    prior = _prior_para(kind, case, note) if kind in ("s138", "recovery", "breach", "termination") else ""
    blocks = list(head)
    blocks.append(Block("subject", text="Sub: " + _final(subject), bold=True))
    blocks.append(Block("p", text="Dear Sir/Madam," if kind == "consumer" else "Sir/Madam,"))
    if opening:
        blocks.append(Block("p", text=_final(opening)))
    for b in body:
        if b.kind == "list":
            blocks.append(b)
            continue
        lay = C.layout(b.text)
        if lay["items"]:
            blocks.append(Block("p", text=_final(lay["lead"] or "The particulars are as follows:")))
            if lay["table"]:
                blocks.append(Block("table", head=lay["table"]["head"], rows=lay["table"]["rows"]))
            else:
                blocks.append(Block("list", items=[_final(x) for x in C.punctuate(lay["items"])]))
            if lay["tail"]:
                blocks.append(Block("p", text=_final(lay["tail"])))
            continue
        b.text = _final(b.text)
        blocks.append(b)
    if li:
        items = []
        for it in li:
            if isinstance(it, tuple):
                items.append((_final(it[0]), it[1]))
                continue
            # Anything that lists things is laid out as points, or as a table
            # with a total when each line carries an amount — never as a
            # paragraph with line breaks buried in it. This applies to every
            # notice type, whatever slot the text came from.
            lay = C.layout(it)
            if lay["items"]:
                if lay["table"]:
                    items.append((_final(lay["lead"] or "The particulars are as follows:"),
                                  lay["table"]))
                else:
                    items.append((_final(lay["lead"] or "The particulars are as follows:"),
                                  dict(sub=[_final(x) for x in C.punctuate(lay["items"])])))
                if lay["tail"]:
                    items.append(_final(lay["tail"]))
            else:
                items.append(_final(it))
        if prior:
            items.append(_final(prior))
        blocks.append(Block("ol", items=items))
    elif prior:
        blocks.append(Block("p", text=_final(prior)))
    for t in tail:
        blocks.append(Block("p", text=_final(t)))
    place = str(case.get("jurisdiction", "")).strip()
    if kind == "recovery" and not place:
        # The recovery reference file: "default to CEAT's registered office unless
        # the user specifies otherwise".
        place = reg_city(ro)
        note(f"No place of jurisdiction was given, so it defaults to the courts at {place} (CEAT's "
             "registered office), as the recovery template directs. Change it if proceedings would be "
             "filed elsewhere.")
    if place and kind in ("s138", "recovery"):
        blocks.append(Block("p", text=_final(_jurisdiction(place))))
    blocks.append(Block("sig", text=sig_block(case)))
    for d in T.drift(kind) + T.drift("clause-library"):
        note(f"Template drift: the skill no longer contains the {d}, so the app's built-in copy was used. "
             "Sync the skill and the app before issuing.")
    return subject, blocks


def blanks(kind: str, case: dict) -> list[str]:
    """Every [● …] still in the notice. None may reach an issued .docx."""
    return C.blanks(to_text(kind, case))


# --------------------------------------------------------------------------
def to_text(kind: str, case: dict) -> str:
    lh = SK.letterhead()
    subject, blocks = build(kind, case)
    out = [lh["name"], lh["address"], lh["meta"], ""]
    for b in blocks:
        if b.kind == "ol":
            for i, it in enumerate(b.items, 1):
                if isinstance(it, tuple):
                    out.append(f"{i}. {it[0]}")
                    t = it[1]
                    if t.get("sub"):
                        for j, sub in enumerate(t["sub"]):
                            out.append(f"   ({chr(97 + j)}) {sub}")
                    else:
                        out.append("   " + " | ".join(t["head"]))
                        for r in t["rows"]:
                            out.append("   " + " | ".join(str(c) for c in r))
                else:
                    out.append(f"{i}. {it}")
                out.append("")
        elif b.kind == "list":
            for j, it in enumerate(b.items):
                out.append(f"   ({chr(97 + j)}) {it}")
            out.append("")
        else:
            out.append(b.text)
            out.append("")
    return "\n".join(out).strip() + "\n"


def to_docx(kind: str, case: dict, specimen: bool = False) -> bytes:
    """The notice as a .docx. `specimen=True` (only on explicit request) marks
    the document itself as a fill-in specimen, as the skill's Step 5a requires
    when blanks are left in."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor, Cm, Inches
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement

    lh = SK.letterhead()
    subject, blocks = build(kind, case)
    doc = docx.Document()

    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(2)
        s.left_margin = s.right_margin = Cm(2)

    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(11)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), "Georgia")

    def para(text, *, size=11, bold=False, underline=False, align=None,
             space_after=8, pre=False):
        p = doc.add_paragraph()
        if align:
            p.alignment = align
        pf = p.paragraph_format
        pf.space_after = Pt(space_after)
        pf.line_spacing = 1.35
        for i, line in enumerate(str(text).split("\n")):
            r = p.add_run(("" if i == 0 else "") + line)
            r.bold = bold
            r.underline = underline
            r.font.size = Pt(size)
            if i < len(str(text).split("\n")) - 1:
                r.add_break()
        return p

    def bottom_border(p, color="9C7A3C", sz=12):
        pPr = p._p.get_or_add_pPr()
        bdr = OxmlElement("w:pBdr")
        bot = OxmlElement("w:bottom")
        bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), str(sz))
        bot.set(qn("w:space"), "4"); bot.set(qn("w:color"), color)
        bdr.append(bot); pPr.append(bdr)

    if specimen:
        sp = para("FILL-IN SPECIMEN — NOT READY TO ISSUE. Complete every [● …] field before use.",
                  size=10, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
        for r in sp.runs:
            r.font.color.rgb = RGBColor(0xB0, 0x1E, 0x1E)

    # letterhead
    para(lh["name"], size=15, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    para(lh["address"], size=9, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    p = para(lh["meta"], size=8, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    bottom_border(p)

    for b in blocks:
        if b.kind == "ol":
            for i, it in enumerate(b.items, 1):
                text, tbl = (it if isinstance(it, tuple) else (it, None))
                p = doc.add_paragraph(style="List Number")
                p.paragraph_format.space_after = Pt(7)
                p.paragraph_format.line_spacing = 1.35
                p.add_run(str(text))
                if tbl and tbl.get("sub"):
                    for j, sub in enumerate(tbl["sub"]):
                        sp = doc.add_paragraph()
                        sp.paragraph_format.left_indent = Inches(0.6)
                        sp.paragraph_format.space_after = Pt(5)
                        sp.paragraph_format.line_spacing = 1.35
                        sp.add_run(f"({chr(97 + j)})\t{sub}")
                elif tbl:
                    t = doc.add_table(rows=1, cols=len(tbl["head"]))
                    t.style = "Table Grid"
                    for c, h in zip(t.rows[0].cells, tbl["head"]):
                        run = c.paragraphs[0].add_run(str(h))
                        run.bold = True
                        run.font.size = Pt(9)
                    for row in tbl["rows"]:
                        cells = t.add_row().cells
                        for c, v in zip(cells, row):
                            r = c.paragraphs[0].add_run(str(v))
                            r.font.size = Pt(9)
                    doc.add_paragraph().paragraph_format.space_after = Pt(4)
        elif b.kind == "list":
            for j, it in enumerate(b.items):
                lp = doc.add_paragraph()
                lp.paragraph_format.left_indent = Inches(0.4)
                lp.paragraph_format.space_after = Pt(5)
                lp.paragraph_format.line_spacing = 1.35
                lp.add_run(f"({chr(97 + j)})\t{it}")
        elif b.kind in ("stamp", "subject"):
            para(b.text, bold=True, underline=b.underline, size=10.5, space_after=4)
        elif b.kind in ("addr", "sig"):
            para(b.text, space_after=14)
        else:
            para(b.text)

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def filename(kind: str, case: dict, specimen: bool = False) -> str:
    import re
    key = {"s138": "Section138_Notice", "recovery": "Recovery_Notice",
           "consumer": "Consumer_Notice_Reply", "breach": "Breach_Notice",
           "termination": "Termination_Notice", "renewal": "Renewal_Notice",
           "fm": "Force_Majeure_Notice", "price": "Price_Revision_Notice"}.get(kind, "Notice")
    party = case.get("client_name") if kind == "consumer" else case.get("noticee_name")
    party = re.sub(r"[^A-Za-z0-9]+", "_", str(party or "Party")).strip("_")[:40] or "Party"
    d = case.get("reply_date") if kind == "consumer" else case.get("notice_date")
    from .words import parse_date
    dt = parse_date(d) or __import__("datetime").date.today()
    return f"CEAT_{key}_{party}_{dt:%Y%m%d}{'_FILL-IN' if specimen else ''}.docx"
