"""One test per defect found when the agent's four notices were reviewed.

Run:  python -m pytest -q tests/
No API key is needed — the model is simulated where the AI path is tested.
"""
import copy
import json
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import pytest                                                    # noqa: E402

from core import analyse as A                                    # noqa: E402
from core import compose as C                                    # noqa: E402
from core import draft as D                                      # noqa: E402
from core import template as T                                   # noqa: E402
from core.validate import validate                               # noqa: E402
from fixtures import (BREACH, BREACH_FIXED, CONSUMER, CONSUMER_FIXED, RECOVERY,  # noqa: E402
                      RECOVERY_FIXED, S138, S138_FIXED)


def text(kind, case):
    return D.to_text(kind, copy.deepcopy(case))


def blockers(kind, case):
    return " | ".join(validate(kind, copy.deepcopy(case)).blockers)


def flags(kind, case):
    return " | ".join(f for _, f in validate(kind, copy.deepcopy(case)).flags)


# ============================================================ Section 138 ===
def test_s138_invoice_total_demand_is_blocked():
    b = blockers("s138", S138)
    assert "4,48,000" in b and "cheque" in b.lower()


def test_s138_no_part_payment_text_is_not_an_explanation():
    """'No further part-payment…' used to downgrade the mismatch to a warning."""
    assert not validate("s138", copy.deepcopy(S138)).ok


def test_s138_analysis_keeps_cheque_amount(monkeypatch):
    """The analysis step overwrote a correct ₹2,00,000 with the ₹4,48,000
    invoice total (_prefer_filtered_rows). Simulate the model and check."""
    case = {k: v for k, v in S138.items() if k not in ("amount", "part_payment")}
    model = A.Found(used_model=True, values={
        "amount": "200000",
        "part_payment": "No further part-payment has been made. Amount demanded is limited to INR 2,00,000/-.",
        "invoices": S138["invoices"], "cheques": S138["cheques"]})
    monkeypatch.setattr(A, "llm_ready", lambda: True)
    monkeypatch.setattr(A, "_call_text", lambda *a, **k: model)
    case["_raw_amt"] = "no part-payment; INR 2,00,000 demanded"
    from core.extract import Doc
    # the trigger: an attached invoice sheet whose total is not the cheque amount
    sheet = Doc(name="invoices.xlsx", kind="tables", digest="t1", tables=[{"sheet": "Sheet1", "rows": [
        ["Invoice No", "Invoice Date", "Customer", "Amount"],
        ["CEAT/CH/2026/0298", "28.05.2026", "Highway Auto Spares Private Limited", "448000"]]}])
    found = A.analyse("s138", case, [sheet])
    assert float(found.values["amount"]) == 200000
    assert "part_payment" not in found.values or not A._NO_PART.search(found.values["part_payment"])


def test_s138_no_drafting_note_paragraph():
    t = text("s138", S138_FIXED)
    assert "No further part-payment" not in t and "limited to" not in t


def test_s138_no_balance_when_nothing_paid():
    t = text("s138", S138_FIXED)
    assert "balance amount" not in t and "payment of the amount against the dishonoured cheque" in t


def test_s138_part_payment_is_stated_as_fact_with_balance():
    c = dict(S138_FIXED, part_paid="50000", part_paid_date="2026-06-25", amount="150000")
    t = text("s138", c)
    assert "paid a sum of INR 50,000/-" in t and "balance of INR 1,50,000/-" in t
    assert validate("s138", c).ok


def test_s138_raw_background_not_pasted():
    t = text("s138", S138)
    assert "3. Invoice No." not in t and "engaged in Dealer" not in t


def test_s138_director_address_not_doubled():
    t = text("s138", S138_FIXED)
    assert "Director, M/s Highway Auto Spares Private Limited, 14," not in t


def test_s138_corrected_passes_with_no_blanks():
    r = validate("s138", copy.deepcopy(S138_FIXED))
    assert r.ok and not r.blanks


# ============================================================== Recovery ===
def test_recovery_not_yet_due_invoices_blocked():
    b = blockers("recovery", RECOVERY)
    assert "not yet due" in b and "3,79,500" in b


def test_recovery_prior_question_not_assumed_no():
    assert "prior notices" in flags("recovery", RECOVERY)


def test_recovery_prior_notice_paragraph_and_overlap_flag():
    t = text("recovery", RECOVERY_FIXED)
    assert "separate notice dated 20.06.2026" in t
    assert "CEAT/CH/2026/0298 are in this statement of account" in flags("recovery", RECOVERY_FIXED)


def test_recovery_interest_start_and_jurisdiction():
    t = text("recovery", RECOVERY_FIXED)
    assert "8% p.a. from the due date of each invoice until realisation." in t
    assert "jurisdiction of the courts at Mumbai" in t
    assert "without a start date" in flags("recovery", RECOVERY)


def test_recovery_corrected_passes():
    r = validate("recovery", copy.deepcopy(RECOVERY_FIXED))
    assert r.ok and not r.blanks


# ============================================================== Consumer ===
def test_consumer_ceat_as_client_blocked():
    assert "CEAT's own" in blockers("consumer", CONSUMER)


def test_consumer_ceat_rejected_from_model_output():
    assert not A._plausible("client_name", "CEAT Limited")
    assert not A._plausible("client_address", "RPG House, 463 Dr. Annie Besant Road")


def test_consumer_subject_and_salutation_at_top():
    t = text("consumer", CONSUMER_FIXED)
    assert t.index("Sub: Reply to your Said Notice.") < t.index("We, CEAT Limited")
    assert t.index("Dear Sir/Madam,") < t.index("We, CEAT Limited")


def test_consumer_no_placeholder_for_missing_inspection():
    t = text("consumer", CONSUMER)
    assert "[● inspection finding]" not in t and "not eligible" not in t


def test_consumer_deny_then_admit_becomes_limited_admission():
    t = text("consumer", CONSUMER_FIXED)
    p2 = [l for l in t.split("\n") if l.startswith("With reference to para 2 ")][0]
    assert "admit only to the limited extent" in p2 and "deny the contents thereof" not in p2


def test_consumer_no_privity_block_not_used_when_dealer_authorised():
    t = text("consumer", CONSUMER)      # para 2 answer calls the dealer "authorised"
    assert "not purchased tyres from us or our authorized dealer" not in t


def test_consumer_contradiction_blocked_if_forced():
    c = dict(CONSUMER_FIXED, dealer_authorised="No")
    assert "Contradiction" in blockers("consumer", c)


def test_consumer_demands_not_denied_twice():
    t = text("consumer", CONSUMER_FIXED)
    assert "The demands raised in the Said Notice" not in t


def test_consumer_same_date_flagged_and_not_seeded():
    assert "same date" in flags("consumer", CONSUMER)


def test_consumer_uses_skill_wording_verbatim():
    t = text("consumer", CONSUMER_FIXED)
    assert "It also appears that your client has not properly apprised you of the facts." in t
    assert "on receipt we issue a claim receipt" in t
    assert "kindly note that we shall have no alternative" in t


def test_consumer_corrected_passes():
    r = validate("consumer", copy.deepcopy(CONSUMER_FIXED))
    assert r.ok and not r.blanks


# ================================================================ Breach ===
def test_breach_sentences_fitted_not_pasted():
    t = text("breach", BREACH)
    for bad in ("obliged to Timely", "in that Under", "constrained to If", "losses., at", ".."):
        assert bad not in t, bad
    assert "Clauses 9.1, 9.2 and 18.1" in t


def test_breach_amounts_indian_format():
    t = text("breach", BREACH)
    # the invoice lines are now a table; amounts are grouped and never raw
    assert "2,84,000" in t and "284000.00" not in t
    assert "INR 7,90,500/-" in t


def test_breach_proprietor_must_be_named():
    assert "proprietor" in blockers("breach", BREACH).lower()


def test_breach_unexplained_clause_flagged():
    assert "18.1" in flags("breach", BREACH)


def test_breach_corrected_passes():
    r = validate("breach", copy.deepcopy(BREACH_FIXED))
    assert r.ok and not r.blanks


# ======================================================= template source ===
def test_wording_comes_from_skill_file(monkeypatch):
    """Editing the skill must change the notice (the README promised this; the
    old code ignored the template text entirely)."""
    orig = T.paragraphs("s138")
    edited = tuple(p.replace("to our shock", "to our surprise") for p in orig)
    monkeypatch.setattr(T, "paragraphs", lambda kind: edited if kind == "s138" else orig)
    assert "to our surprise" in text("s138", S138_FIXED)


def test_missing_template_paragraph_reported_as_drift(monkeypatch):
    orig = T.paragraphs("s138")
    trimmed = tuple(p for p in orig if not p.startswith("That in the course of business"))
    monkeypatch.setattr(T, "paragraphs", lambda kind: trimmed if kind == "s138" else orig)
    assert "Template drift" in flags("s138", S138_FIXED)


# ============================================================ blanks gate ===
def test_blanks_reported():
    c = dict(S138_FIXED)
    c.pop("memo_date")
    c["cheques"] = [dict(S138_FIXED["cheques"][0], bank="")]
    r = validate("s138", c)
    assert any("drawn-on bank" in b for b in r.blanks)


def test_specimen_docx_is_marked():
    import io
    import docx
    raw = D.to_docx("s138", copy.deepcopy(S138_FIXED), specimen=True)
    first = docx.Document(io.BytesIO(raw)).paragraphs[0].text
    assert first.startswith("FILL-IN SPECIMEN")


# ======================================================== the AI paths ===
def test_ai_fit_accepted_when_faithful():
    case = copy.deepcopy(BREACH_FIXED)
    reply = json.dumps({"mode": "inline", "text": "make payment for all Products within 30 (thirty) days "
                                                  "of each invoice and pay interest at 8% per annum on "
                                                  "overdue amounts"})
    C.llm_fit_case("breach", case, chat=lambda m: reply)
    t = text("breach", case)
    assert "you were obliged to make payment for all Products within 30 (thirty) days" in t


def test_ai_fit_rejected_when_it_adds_a_fact():
    case = copy.deepcopy(BREACH_FIXED)
    reply = json.dumps({"mode": "inline", "text": "pay INR 9,99,999/- by 01.01.2027"})
    notes = C.llm_fit_case("breach", case, chat=lambda m: reply)
    assert any("rejected" in n for n in notes)
    assert "9,99,999" not in text("breach", case)


def test_ai_fit_goes_stale_when_answer_changes():
    case = copy.deepcopy(BREACH_FIXED)
    C.llm_fit_case("breach", case, chat=lambda m: json.dumps(
        {"mode": "inline", "text": "make payment within 30 (thirty) days"}))
    case["obligation"] = "maintain minimum stock of 400 tyres"
    assert "maintain minimum stock of 400 tyres" in text("breach", case)


def test_ai_review_flags_are_collected():
    reply = json.dumps({"issues": [{"severity": "high", "where": "para 10",
                                    "problem": "Demand differs from cheque amount"}]})
    out = C.ai_review("s138", "notice text", {"amount": "1"}, chat=lambda m: reply)
    assert out == [("crit", "AI read-through (para 10): Demand differs from cheque amount")]


def test_chat_retries_without_temperature(monkeypatch):
    """gpt-5.x rejects temperature=0; that used to kill every model call."""
    calls = []

    class Completions:
        def create(self, **kw):
            calls.append(kw)
            if "temperature" in kw:
                raise Exception("Unsupported value: 'temperature' does not support 0 with this model.")
            return "ok"

    class Client:
        chat = type("X", (), {"completions": Completions()})()

    monkeypatch.setattr("core.models.pick", lambda role: ("gpt-5.1", ""))
    assert A._chat(Client(), "text", temperature=0, messages=[]) == "ok"
    assert "temperature" not in calls[-1]


def test_openai_never_falls_back_to_arbitrary_model(monkeypatch):
    from core import models as M
    monkeypatch.setattr(M, "llm_settings", lambda: dict(key="k", base="", text="gpt-5.1", vision="gpt-5.1",
                                                        name="OpenAI-compatible"))
    monkeypatch.setattr(M, "available", lambda force=False: {"babbage-002", "gpt-4o"})
    assert M.pick("text")[0] == "gpt-4o"
    monkeypatch.setattr(M, "available", lambda force=False: {"babbage-002", "davinci-002"})
    assert M.pick("text")[0] == "gpt-5.1"


# ======================================================== freetext parse ===
@pytest.mark.parametrize("answer,paid,amount", [
    ("No further part-payment has been made. Amount demanded is limited to INR 2,00,000/-.", "0", "200000"),
    ("INR 50,000 received on 25.06.2026; balance INR 1,50,000 demanded.", "50000", "150000"),
])
def test_part_payment_answer_parsed(answer, paid, amount):
    from core.freetext import _parse
    v = _parse("s138", "amt", answer, {})
    assert v["part_paid"] == paid and float(v["amount"]) == float(amount)


def test_default_chat_path_through_fake_client(monkeypatch):
    """The real wiring: compose -> analyse._client/_chat -> client.chat.completions."""
    class Msg:
        content = json.dumps({"mode": "inline", "text": "terminate the Agreement for cause under Clause 19.1"})

    class Rsp:
        choices = [type("Ch", (), {"message": Msg()})()]

    class Completions:
        def create(self, **kw):
            assert kw["response_format"] == {"type": "json_object"}
            return Rsp()

    class Client:
        chat = type("X", (), {"completions": Completions()})()

    monkeypatch.setattr(A, "_client", lambda: (Client(), {}))
    monkeypatch.setattr("core.models.pick", lambda role: ("gpt-5.1", ""))
    case = copy.deepcopy(BREACH_FIXED)
    case["consequences"] = "terminate the Dealership Agreement for cause under Clause 19.1"
    C.llm_fit_case("breach", case)
    assert case["_fit"]["consequences"]["text"].startswith("terminate the Agreement")
    assert "constrained to terminate the Agreement for cause under Clause 19.1" in text("breach", case)


# ================================ unseen cases, typed as a user would ===
def _typed(kind, case):
    """The app's own path with no model key: analyse, apply, validate."""
    from core.schema import LABELS, TABLE_COLS, is_filled
    case = copy.deepcopy(case)
    found = A.analyse(kind, case, [])
    for k, v in (found.values or {}).items():
        if (k in LABELS or k in TABLE_COLS) and not is_filled(case, k) and v not in (None, "", [], {}):
            case[k] = v
    return case, validate(kind, case)


NEW_S138 = {
    "noticee_type": "Partnership + partners", "_defaults": [],
    "_raw_addr": "M/s Sai Krishna Tyre Traders, a partnership firm, partners Mr. Venkat Rao and Mrs. Lakshmi "
                 "Rao, 5-2-110 Ranigunj, Secunderabad – 500003",
    "background": "They have been a CEAT dealer in Secunderabad since 2019, buying on 45-day credit.",
    "cheques": [{"no": "118822", "date": "2026-07-10", "amt": "150000", "bank": "State Bank of India, Ranigunj"},
                {"no": "118823", "date": "2026-07-20", "amt": "125000", "bank": "State Bank of India, Ranigunj"}],
    "_raw_bounce": "Both cheques were presented on 25.07.2026 and returned on 28.07.2026 with bank memo of the "
                   "same date, reason 'Payment stopped by drawer'.",
    "_raw_inv": "INV/HYD/2026/771 dated 02.06.2026 for Rs. 3,10,000",
    "_raw_amt": "They paid Rs. 25,000 on 05.08.2026 after the cheques bounced. Balance Rs. 2,50,000 is demanded.",
    "jurisdiction": "Hyderabad", "_raw_prior": "No earlier notice.",
    "_raw_sig": "Meena Marar, General Manager – Legal, authorised to sign", "notice_date": "2026-08-12",
}

NEW_BREACH = {
    "noticee_type": "Individual / sole proprietor", "_defaults": [],
    "_raw_addr": "Shree Balaji Tyres, owned by Mr. Ramesh Patil, Shop 4, Station Road, Nashik – 422001",
    "_raw_agr": "Distributorship Agreement dated 15.01.2025, clause 7.3",
    "obligation": "lift a minimum of 500 tyres every quarter",
    "breach_facts": "you lifted only 180 tyres in the quarter ended 30.06.2026 against the minimum of 500",
    "cure_period": "30 (THIRTY) days",
    "consequences": "termination of the Distributorship Agreement under Clause 14.2",
    "_raw_prior": "no", "_raw_sig": "Meena Marar, General Manager – Legal, authorised", "notice_date": "2026-09-22",
}


def test_unseen_cheque_case_read_without_a_key(monkeypatch):
    monkeypatch.setattr(A, "llm_ready", lambda: False)
    case, r = _typed("s138", NEW_S138)
    assert case["noticee_name"] == "M/s Sai Krishna Tyre Traders"
    assert [d["name"] for d in case["directors"]] == ["Mr. Venkat Rao", "Mrs. Lakshmi Rao"]
    assert case["invoices"][0] == {"no": "INV/HYD/2026/771", "date": "2026-06-02", "amt": "310000"}
    assert case["memo_date"] == "2026-07-28" and case["dishonour_reason"] == "Payment stopped by drawer"
    assert float(case["amount"]) == 250000 and float(case["part_paid"]) == 25000
    assert r.ok and not r.blanks
    t = D.to_text("s138", case)
    assert "That Noticee No. 1 has been a CEAT dealer" in t
    assert "balance of INR 2,50,000/-" in t


def test_unseen_breach_case_read_without_a_key(monkeypatch):
    monkeypatch.setattr(A, "llm_ready", lambda: False)
    case, r = _typed("breach", NEW_BREACH)
    assert case["noticee_name"] == "Mr. Ramesh Patil" and case["firm_name"] == "Shree Balaji Tyres"
    assert case["agreement_name"] == "Distributorship Agreement" and case["clause_no"] == "7.3"
    assert r.ok and not r.blanks
    t = D.to_text("breach", case)
    assert "you were obliged to lift a minimum of 500 tyres every quarter." in t
    assert "in that you lifted only 180 tyres" in t
    assert "constrained to proceed with termination of the Distributorship Agreement" in t


@pytest.mark.parametrize("answer", [
    "INV/HYD/2026/771 dated 02.06.2026 for Rs. 3,10,000",
    "INV/HYD/2026/771 | 02.06.2026 | 3,10,000",
    "INV/HYD/2026/771  02.06.2026  310000",
    "Invoice No. INV/HYD/2026/771 dated 02.06.2026 for INR 3,10,000/-",
])
def test_invoice_formats_never_split_indian_amounts(answer):
    from core.freetext import _parse
    assert _parse("s138", "inv", answer, {})["invoices"][0]["amt"] == "310000"


@pytest.mark.parametrize("answer,name,firm", [
    ("Shree Balaji Tyres owned by Mr. Ramesh Patil, Shop 4, Nashik", "Mr. Ramesh Patil", "Shree Balaji Tyres"),
    ("Sharma Tyres (Prop. Mr. Rakesh Sharma), Shop 14, MG Road, Jaipur", "Mr. Rakesh Sharma", "Sharma Tyres"),
    ("M/s Deccan Auto Tyres, Prop. Mr. Sanjay Kulkarni, Plot 12, Pune", "Mr. Sanjay Kulkarni", "M/s Deccan Auto Tyres"),
])
def test_proprietor_phrasings(answer, name, firm):
    from core.freetext import _parse
    v = _parse("breach", "addr", answer, {"noticee_type": "Individual / sole proprietor"})
    assert v["noticee_name"] == name and v["firm_name"] == firm


# ================== the contract notices (termination / renewal / fm / price) ===
from contract_fixtures import FM, PRICE, REN, TERM                      # noqa: E402


def test_price_table_typed_by_hand_is_read():
    """The price notice could not be drafted at all: the row reader only filled a
    column called "amt", so a price table (sku|old|nw|pct) lost every figure."""
    from core.freetext import _parse
    typed = ("SKU | Tyre size | Existing price (INR) | Revised price (INR) | Change\n"
             "TBR-2958 | 295/80 R22.5 | 24,500 | 26,020 | +6.2%\n"
             "LCV-0750 | 7.50-16 | 6,450 | 6,850 | +6.2%")
    rows_ = _parse("price", "prices", typed, {})["prices"]
    assert len(rows_) == 2
    assert rows_[0] == {"sku": "TBR-2958 295/80 R22.5", "old": "24500", "nw": "26020", "pct": "+6.2%"}


def test_typed_answers_reach_the_model_prompt():
    """A table typed into a box was never sent to the model — so the model
    reported, correctly, that no price table had been supplied."""
    prompt = A._prompt("price", {"_raw_prices": "TBR-2958 | 24,500 | 26,020 | +6.2%"}, [])
    assert "TBR-2958 | 24,500 | 26,020" in prompt


def test_price_notice_drafts_and_states_percentages():
    r = validate("price", copy.deepcopy(PRICE))
    assert r.ok and not r.blanks
    t = text("price", PRICE)
    assert "26,020" in t and "+6.2%" in t and "The revised prices are as follows:" in t


def test_price_notice_accepts_an_overall_percentage_only():
    c = dict(PRICE, prices=[], price_change_pct="6.2")
    r = validate("price", c)
    assert r.ok
    assert "by 6.2%," in text("price", c)


def test_price_notice_without_either_is_blocked():
    c = dict(PRICE, prices=[], price_change_pct="")
    assert "Revised prices" in blockers("price", c)


def test_termination_separates_obligation_and_termination_clauses():
    t = text("termination", TERM)
    assert "In terms of Clause 7.2, you were required to register" in t
    assert "termination of the Agreement under Clause 12.1(a)" in t
    assert "required to to" not in t


def test_termination_warns_when_one_clause_does_both():
    c = dict(TERM, obligation_clause="12.1(a)")
    assert "both cited to Clause 12.1(a)" in flags("termination", c)


def test_termination_wind_down_is_a_lettered_list_without_duplicate_dues():
    t = text("termination", TERM)
    assert "(a) pay the balance" in t and "(e) retain every tyre" in t
    assert t.count("Upon termination, you are called upon") == 1
    assert "clear all outstanding dues of INR 1,24,400/- and to:" not in t


def test_termination_blocks_retrospective_effect():
    assert "retrospectively" in blockers("termination", dict(TERM, effective_date="2026-09-01"))


def test_fm_sentences_are_fitted_not_pasted():
    t = text("fm", FM)
    for bad in ("impact/duration is due to", "and to suspension of", "performing CEAT"):
        assert bad not in t, bad
    assert "The expected impact and duration is that production" in t
    assert "(a) 120 units" in t and "(e) the Company will send" in t


def test_fm_notice_speaks_of_itself_as_the_company():
    t = text("fm", FM)
    body = t[t.index("Sir/Madam"):]
    assert "CEAT will" not in body and "CEAT's" not in body


def test_fm_flags_the_notice_deadline():
    f = flags("fm", FM)
    assert "8 day(s) later" in f and "notice deadline" in f


def test_renewal_keeps_the_firm_as_noticee_one():
    t = text("renewal", REN)
    assert "Noticee No. 1\nM/s Om Sai Tyres & Services" in t
    assert "Partner, M/s Om Sai Tyres & Services" in t


def test_renewal_terms_are_itemised_with_rupee_amounts():
    t = text("renewal", REN)
    assert "(c) Security deposit" in t
    assert "INR 2,00,000/- (Rupees Two Lakh Only)" in t
    assert "INR 1,80,00,000/- (Rupees One Crore Eighty Lakh Only)" in t
    assert "200000" not in t and "18000000" not in t


def test_renewal_states_what_happens_if_not_accepted():
    t = text("renewal", REN)
    assert "if your written acceptance is not received" in t and "expire on 31.12.2026" in t


def test_renewal_confirm_by_must_precede_expiry():
    assert "Bring the confirm-by date forward" in blockers("renewal", dict(REN, confirm_by="2027-01-15"))


def test_non_renewal_time_bar_is_blocked():
    # 100 days remain to expiry, so a 120-day notice requirement cannot be met
    c = dict(REN, intent="Do not renew", notice_period="120 days", revised_terms="", confirm_by="")
    assert "notice of non-renewal" in blockers("renewal", c)


def test_contact_person_is_not_made_a_noticee():
    c = dict(FM, noticee_type="Company + directors",
             directors=[{"name": "Mr. Rahul Sawant", "capacity": "Director – Operations", "address": ""}])
    t = text("fm", c)
    assert "Noticee No. 2" not in t
    assert "Kind Attn.: Mr. Rahul Sawant, Director – Operations" in t


def test_firm_put_back_as_noticee_when_model_returns_people():
    f = A.Found(values={"noticee_type": "Partnership + partners",
                        "noticee_name": "Harish Patel and Nilesh Patel",
                        "firm_name": "M/s Om Sai Tyres & Services"})
    A._firm_first({}, f)
    assert f.values["noticee_name"] == "M/s Om Sai Tyres & Services"


def test_bare_amounts_get_house_formatting_but_quantities_do_not():
    m = C.money_full
    assert m("deposit of 200000") == "deposit of INR 2,00,000/- (Rupees Two Lakh Only)"
    assert m("deliver 400 tyres per month") == "deliver 400 tyres per month"
    assert m("1800 tyres lifted") == "1800 tyres lifted"
    assert m("interest at 12% per annum") == "interest at 12% per annum"


def test_excel_formulas_without_saved_values_are_reported():
    import io as _io
    from openpyxl import Workbook
    from core.extract import extract
    wb = Workbook()
    wb.active.append(["SKU", "Old", "Revised"])
    wb.active.append(["TBR-2958", 24500, "=B2*1.062"])
    buf = _io.BytesIO()
    wb.save(buf)
    note = extract("Price Working.xlsx", buf.getvalue()).note
    assert "no saved result" in note and "read as EMPTY" in note


def test_all_four_contract_notices_draft_clean():
    for kind, case in (("termination", TERM), ("fm", FM), ("renewal", REN), ("price", PRICE)):
        r = validate(kind, copy.deepcopy(case))
        assert r.ok and not r.blanks, (kind, r.blockers, r.blanks)
        assert not [f for _, f in r.flags if f.startswith("Read-through")], kind


def test_fm_quantities_table_is_read_and_totalled():
    from core.freetext import _parse
    rows_ = _parse("fm", "qty", "September 2026 | 400 | 160 | 240 | 25.09.2026\n"
                                "October 2026 | 400 | 0 | 400 | 25.10.2026", {})["quantities"]
    assert rows_[0] == {"period": "September 2026", "sched": "400", "done": "160", "pend": "240",
                        "due": "2026-09-25"}
    t = text("fm", FM)
    assert "Month / period | Scheduled | Delivered | Pending | Due date" in t
    assert "Total | 800 | 160 | 640" in t


def test_fm_without_a_quantities_table_is_prompted_for_one():
    assert "No quantities table" in flags("fm", dict(FM, quantities=[]))


def test_termination_prompts_for_a_schedule_of_particulars():
    assert "schedule of particulars" in flags("termination", TERM)
    assert "schedule of particulars" not in flags(
        "termination", dict(TERM, extra_notes="Schedule of the 19 claims is annexed."))


def test_pending_is_worked_out_when_not_given():
    from core.freetext import _parse
    r = _parse("fm", "qty", "September 2026 | 400 | 160 | | 25.09.2026", {})["quantities"][0]
    assert r["pend"] == "240"



# ======================= points and tables instead of prose ===============
def test_listed_invoices_become_a_table_with_a_total():
    t = text("breach", BREACH)
    assert "Invoice / reference | Date | Amount (INR)" in t
    assert "CEAT/PN/2026/0451 | 18.05.2026 | 2,84,000" in t
    assert "Total |  | 7,90,500" in t


def test_a_total_line_is_not_swallowed_into_the_table():
    t = text("breach", BREACH)
    assert "The total outstanding on your account is INR 7,90,500/-" in t


def test_plain_multi_item_answers_become_lettered_points():
    lay = C.layout("Upon termination you must:\n- remove all signage within 7 days\n"
                   "- return the dealership certificate\n- cease use of the trademarks")
    assert lay["table"] is None and len(lay["items"]) == 3
    t = text("termination", TERM)
    assert "(a) pay the balance" in t and "(c) cease all use" in t


def test_consumer_demands_are_listed_not_run_together():
    c = dict(CONSUMER_FIXED)
    c["paras"] = [dict(p, text=p["text"].replace("The demands for", "The claims for"))
                  for p in c["paras"]]
    t = text("consumer", c)
    assert "The demands raised in the Said Notice, namely:" in t
    assert "(a) Replace four tyres free of cost or refund INR 32,000/-;" in t
    assert "(d) Pay INR 20,000/- towards cost of legal notice," in t
    assert "are each wholly untenable" in t


def test_amounts_in_reply_paragraphs_are_formatted():
    t = text("consumer", CONSUMER_FIXED)
    assert "for INR 30,000/-" in t and "for 30000" not in t


def test_years_quantities_and_distances_are_not_turned_into_money():
    m = C.money_full
    assert m("valid for 2026 and 2027", words=False) == "valid for 2026 and 2027"
    assert m("tread life of 40000 km", words=False) == "tread life of 40000 km"
    assert m("for 400 units", words=False) == "for 400 units"
