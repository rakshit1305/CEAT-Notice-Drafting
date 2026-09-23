"""Turn an uploaded file into something the analyser can reason over.

Spreadsheets come back as tables (every sheet), documents as text, images as
base64 for the vision model. Nothing is interpreted here — that is analyse.py.
"""
from __future__ import annotations
import base64
import hashlib
import io
import re
from dataclasses import dataclass, field

from .config import MAX_DOC_CHARS, MAX_SHEET_ROWS

IMAGE_EXT = {"jpg", "jpeg", "png", "webp", "gif", "bmp", "tif", "tiff"}
SHEET_EXT = {"xlsx", "xls", "xlsm", "csv"}
DOC_EXT = {"pdf", "docx", "doc", "txt", "md"}


@dataclass
class Doc:
    name: str
    kind: str = "text"                 # text | tables | image | error
    text: str = ""
    tables: list[dict] = field(default_factory=list)   # {sheet, rows:[[...]]}
    b64: str = ""
    mime: str = ""
    note: str = ""
    digest: str = ""

    @property
    def dropped_rows(self) -> int:
        """Rows in the file that MAX_SHEET_ROWS kept from the model.

        validate.py already blocks on this; without the property it read 0
        through getattr and a truncated ledger passed silently.
        """
        if self.kind != "tables":
            return 0
        return sum(max(0, len(t["rows"]) - MAX_SHEET_ROWS) for t in self.tables)

    @property
    def dropped_chars(self) -> int:
        if self.kind == "tables":
            full = 0
            for t in self.tables:
                for r in t["rows"]:
                    full += len(" | ".join("" if c is None else str(c).strip() for c in r)) + 1
        else:
            full = len(self.text or "")
        return max(0, full - MAX_DOC_CHARS)

    def as_prompt(self) -> str:
        """A compact, faithful rendering for the model."""
        if self.kind == "tables":
            out = []
            for t in self.tables:
                out.append(f"--- sheet: {t['sheet']} ({len(t['rows'])} rows) ---")
                for r in t["rows"][:MAX_SHEET_ROWS]:
                    cells = ["" if c is None else str(c).strip() for c in r]
                    if any(cells):
                        out.append(" | ".join(cells))
                if len(t["rows"]) > MAX_SHEET_ROWS:
                    out.append(f"... {len(t['rows']) - MAX_SHEET_ROWS} further rows not shown")
            return "\n".join(out)[:MAX_DOC_CHARS]
        return (self.text or "")[:MAX_DOC_CHARS]


def _digest(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()[:16]


def extract(name: str, data: bytes) -> Doc:
    ext = (name.rsplit(".", 1)[-1] if "." in name else "").lower()
    d = Doc(name=name, digest=_digest(data))

    if ext in IMAGE_EXT:
        d.kind = "image"
        d.b64 = base64.b64encode(data).decode()
        d.mime = f"image/{'jpeg' if ext in ('jpg', 'jpeg') else ext}"
        return d

    try:
        if ext == "csv":
            d.kind = "tables"
            d.tables = [dict(sheet="csv", rows=_csv_rows(data))]
        elif ext in ("xlsx", "xls", "xlsm"):
            d.kind = "tables"
            d.tables = _excel_tables(data)
            d.note = _formula_note(data, ext, d.tables) or d.note
        elif ext == "pdf":
            d.text, pages, has_text = _pdf_text(data)
            if not has_text:
                d.kind = "error"
                d.note = (f"no text layer across {pages} page(s) — this is a scan, so it is an image. "
                          "Export it as an image and attach that, and the vision model will read it.")
            else:
                d.kind = "text"
        elif ext == "docx":
            d.kind = "text"
            d.text = _docx_text(data)
        elif ext == "doc":
            d.kind = "error"
            d.note = "legacy .doc cannot be parsed — save as .docx or PDF"
        else:
            d.kind = "text"
            d.text = data.decode("utf-8", errors="replace")
    except Exception as e:                                  # never crash the app on a bad file
        d.kind = "error"
        d.note = f"{type(e).__name__}: {e}"
    return d


# --------------------------------------------------------------------------
def _csv_rows(data: bytes) -> list[list]:
    import csv
    text = data.decode("utf-8-sig", errors="replace")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except Exception:
        dialect = csv.excel
    rows = [r for r in csv.reader(io.StringIO(text), dialect)]
    return [r for r in rows if any(str(c).strip() for c in r)]


def _excel_tables(data: bytes) -> list[dict]:
    import pandas as pd
    book = pd.read_excel(io.BytesIO(data), sheet_name=None, header=None, dtype=object)
    out = []
    for sheet, df in book.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        rows = [["" if (v is None or (isinstance(v, float) and v != v)) else v for v in row]
                for row in df.values.tolist()]
        if rows:
            out.append(dict(sheet=str(sheet), rows=rows))
    return out


def _formula_note(data: bytes, ext: str, tables: list[dict]) -> str:
    """A sheet written by a script holds formulas with no saved result. Excel
    files are read for their saved values, so those cells arrive empty — which
    looks to everyone like the app simply ignored the figures."""
    if ext != "xlsx":
        return ""
    try:
        from openpyxl import load_workbook
        raw = load_workbook(io.BytesIO(data), data_only=False, read_only=True)
        cached = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
        formulas = empty = 0
        for ws in raw.worksheets:
            cs = cached[ws.title]
            for row in ws.iter_rows():
                for c in row:
                    if isinstance(c.value, str) and c.value.startswith("="):
                        formulas += 1
                        if cs.cell(c.row, c.column).value in (None, ""):
                            empty += 1
        raw.close()
        cached.close()
    except Exception:
        return ""
    if not formulas:
        return ""
    if empty:
        return (f"{empty} of this workbook's {formulas} formula cell(s) have no saved result, so those "
                "figures read as EMPTY here. Open the file in Excel and save it, or type the figures "
                "into the answer box — otherwise they cannot reach the notice.")
    return f"{formulas} cell(s) are formulas; their saved results were read."


def _pdf_text(data: bytes) -> tuple[str, int, bool]:
    import pdfplumber
    chunks: list[str] = []
    pages = 0
    with pdfplumber.open(io.BytesIO(data)) as pdf:
        pages = len(pdf.pages)
        for page in pdf.pages[:30]:
            chunks.append(page.extract_text() or "")
            for tbl in (page.extract_tables() or []):
                for row in tbl:
                    cells = [("" if c is None else str(c).strip()) for c in row]
                    if any(cells):
                        chunks.append(" | ".join(cells))
    text = re.sub(r"\n{3,}", "\n\n", "\n".join(chunks)).strip()
    return text, pages, len(re.sub(r"\s", "", text)) >= 40


def _docx_text(data: bytes) -> str:
    import docx
    doc = docx.Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return re.sub(r"\n{3,}", "\n\n", "\n".join(parts)).strip()
